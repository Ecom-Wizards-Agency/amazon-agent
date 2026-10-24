import importlib.util,tempfile,unittest
from pathlib import Path
from openpyxl import Workbook
spec=importlib.util.spec_from_file_location('builder',Path(__file__).resolve().parents[1] / 'build_keyword_workbook.py');b=importlib.util.module_from_spec(spec);spec.loader.exec_module(b)
MASTER=[['Search Terms','SV','Relev.','B087654321'],['iron',1000,1,1],['anemia',500,1,2]]
EXPANDED=[['Search Terms','SV','Relev.','B087654321'],['anemia',500,1,2],['anemia pills',100,.6,3]]
class KeywordSourceLabelTests(unittest.TestCase):
 def cfg(self,**extra):
  c={'product_anchor':{'asin':'','marketplace':'US','client':'Example'},'ranking_context':{'mode':'unavailable'},'triage':{'brand_tokens':[],'form_tokens':[],'claim_tokens':['anemia'],'negative_tokens':[],'match':'word','categories_order':[],'semantic_opportunity':{'min_relevancy':.3,'min_search_volume':1,'include_if_anchor_not_ranking':False}},'final_action':{'by_classification':{'Unsupported claim/health-risk':'Ignore'},'allowed':['Ignore']}}
  c.update(extra);return c
 def outlier(self,cfg):
  ws=Workbook().active;b.build_outlier(ws,MASTER,EXPANDED,set(),cfg,{},[]);return list(ws.values)
 def test_label_helper_defaults(self):
  for c in ({},{'keyword_source':''},{'keyword_source':'   '},{'keyword_source':None}):self.assertEqual(b.keyword_source_label(c),'DataDive')
  self.assertEqual(b.keyword_source_label({'keyword_source':' SoldScope '}),'SoldScope')
 def test_default_outlier_labels_unchanged(self):
  rows=self.outlier(self.cfg());self.assertEqual(rows[0][2],'DataDive Relevancy')
  sources={r[0]:r[10] for r in rows[1:]};self.assertEqual(sources,{'anemia':'DataDive Master List','anemia pills':'DataDive Expanded MKL 1%'})
 def test_soldscope_outlier_labels(self):
  rows=self.outlier(self.cfg(keyword_source='SoldScope'));self.assertEqual(rows[0][2],'SoldScope Relevancy')
  sources={r[0]:r[10] for r in rows[1:]};self.assertEqual(sources,{'anemia':'SoldScope Master List','anemia pills':'SoldScope Expanded MKL 1%'})
  self.assertFalse(any('DataDive' in str(v) for r in rows for v in r if v is not None))
 def wb(self):
  wb=Workbook();wb.active.title='Overview';wb.create_sheet('3.1 MKL DataDive 30%');wb.create_sheet('2.1 MKL DataDive 1%');wb.create_sheet('Notes');return wb
 def test_template_renames_in_place(self):
  wb=self.wb();w=[];ws30=wb['3.1 MKL DataDive 30%'];ws1=wb['2.1 MKL DataDive 1%']
  a=b.apply_template_renames(wb,{'tabs':{'template_renames':{'3.1 MKL DataDive 30%':'3.1 MKL SoldScope 30%','2.1 MKL DataDive 1%':'2.1 MKL SoldScope 1%','Notes':'Notes'}}},w)
  self.assertEqual(wb.sheetnames,['Overview','3.1 MKL SoldScope 30%','2.1 MKL SoldScope 1%','Notes']);self.assertEqual(w,[])
  self.assertIs(wb['3.1 MKL SoldScope 30%'],ws30);self.assertIs(wb['2.1 MKL SoldScope 1%'],ws1)
  self.assertEqual(a,[{'from':'3.1 MKL DataDive 30%','to':'3.1 MKL SoldScope 30%'},{'from':'2.1 MKL DataDive 1%','to':'2.1 MKL SoldScope 1%'}])
 def test_template_rename_missing_old_warns(self):
  wb=self.wb();w=[];a=b.apply_template_renames(wb,{'tabs':{'template_renames':{'Nope':'X'}}},w)
  self.assertEqual(a,[]);self.assertEqual(len(w),1);self.assertIn('skipped',w[0]);self.assertIn('Nope',w[0])
  self.assertEqual(wb.sheetnames,['Overview','3.1 MKL DataDive 30%','2.1 MKL DataDive 1%','Notes'])
 def test_template_rename_collision_refused(self):
  wb=self.wb();w=[];a=b.apply_template_renames(wb,{'tabs':{'template_renames':{'3.1 MKL DataDive 30%':'Notes'}}},w)
  self.assertEqual(a,[]);self.assertEqual(len(w),1);self.assertIn('refused',w[0]);self.assertEqual(wb.sheetnames,['Overview','3.1 MKL DataDive 30%','2.1 MKL DataDive 1%','Notes'])
 def test_stale_rename_references_warn(self):
  applied=[{'from':'3.1 MKL DataDive 30%','to':'3.1 MKL SoldScope 30%'}]
  cfg={'tabs':{'exact_paste':{'3.1 MKL DataDive 30%':'master_csv'},'rebuild':[]},'tab_order':['Overview','3.1 MKL DataDive 30%'],'stale_data_guard':{'tabs':['3.1 MKL SoldScope 30%']}}
  w=b.stale_rename_references(cfg,applied);self.assertEqual(len(w),2)
  self.assertTrue(any('tabs.exact_paste' in m for m in w));self.assertTrue(any('tab_order' in m for m in w))
  cfg_ok={'tabs':{'exact_paste':{'3.1 MKL SoldScope 30%':'master_csv'}},'tab_order':['3.1 MKL SoldScope 30%']}
  self.assertEqual(b.stale_rename_references(cfg_ok,applied),[]);self.assertEqual(b.stale_rename_references(cfg,[]),[])
 def test_no_renames_configured(self):
  wb=self.wb();w=[];self.assertEqual(b.apply_template_renames(wb,{},w),[]);self.assertEqual(w,[])
 def handoff(self,cfg):
  with tempfile.TemporaryDirectory() as d:
   p=str(Path(d)/'note.md');m={'outputs':{'workbook_local':'w.xlsx','manifest':'m.json'}};b.write_handoff_note(p,cfg,{},m);return Path(p).read_text()
 def test_handoff_heading_uses_label(self):
  self.assertIn('\n## DataDive Exports\n',self.handoff(self.cfg()))
  t=self.handoff(self.cfg(keyword_source='SoldScope'));self.assertIn('\n## SoldScope Exports\n',t);self.assertNotIn('DataDive',t)
if __name__=='__main__':unittest.main()
