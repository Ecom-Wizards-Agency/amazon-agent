import copy
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

PATH = Path(__file__).resolve().parents[1] / 'operations.py'
spec = importlib.util.spec_from_file_location('operations', PATH)
op = importlib.util.module_from_spec(spec)
spec.loader.exec_module(op)

class OperationsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.service = op.Operations(self.root / 'state')
        self.account = {'client_slug':'test-brand','profile_key':'test-us','marketplace':'US'}
        self.source = self.root / 'source.csv'
        self.source.write_text('sku,item_name.0.value,bullet_point.0.value\na,Old A,Keep A\nb,Old B,Keep B\n')
        self.request = {'schema_version':1,'operation_id':'run-1','operation':'seo.apply','account':self.account,'targets':['a','b'],'inputs':{
            'live':{'account':self.account,'observed_at':op.now(),'rows':{'a':{'item_name.0.value':'Old A','bullet_point.0.value':'Keep A'},'b':{'item_name.0.value':'Old B','bullet_point.0.value':'Keep B'}}},
            'approved_seo':{'account':self.account,'source_id':'seo-revision-1','approved':True,'approved_by':'operator','approved_at':op.now(),'rows':{'a':{'item_name.0.value':'New A'},'b':{'bullet_point.0.value':'New B'}}},
            'source_export':{'path':str(self.source),'sha256':op.file_hash(self.source)},'scope_fields':['item_name.0.value','bullet_point.0.value']}}
    def prepare(self):
        result=self.service.prepare(self.request)
        self.plan=json.loads(Path(result['plan_path']).read_text())
        return result
    def execute_request(self,result):
        return {'schema_version':1,'operation_id':'run-1','plan_hash':result['plan_hash'],'grant':{'execute':True,'account':self.account,'operation':'seo.apply','targets':['a','b'],'plan_hash':result['plan_hash'],'allow_live_canary':True}}
    def evidence_request(self,result):
        return {'schema_version':1,'operation_id':'run-1','plan_hash':result['plan_hash'],'evidence':{'account':self.account,'plan_hash':result['plan_hash'],'source_id':'amazon-feed-and-listing','observed_at':op.now(),'submission_id':'feed-1','processing_status':'complete','rows':copy.deepcopy(self.plan['body'].get('expected_rows', {}))}}
    def assertCode(self,code,fn):
        with self.assertRaises(op.OperationError) as raised:fn()
        self.assertEqual(raised.exception.code,code)
    def test_full_grid_preserves_unchanged_values(self):
        self.prepare()
        self.assertEqual(self.plan['body']['expected_rows'],{'a':{'item_name.0.value':'New A','bullet_point.0.value':'Keep A'},'b':{'item_name.0.value':'Old B','bullet_point.0.value':'New B'}})
        self.assertTrue(Path(self.plan['body']['upload']).is_file())
    def test_idempotent_prepare_and_immutable_request(self):
        first=self.prepare();self.assertEqual(self.service.prepare(self.request)['plan_hash'],first['plan_hash'])
        self.request['inputs']['approved_seo']['rows']['a']['item_name.0.value']='Different'
        self.assertCode('immutable_request',lambda:self.service.prepare(self.request))
    def test_cross_account_input_rejected(self):
        self.request['inputs']['live']['account']=dict(self.account,marketplace='DE')
        self.assertCode('identity_mismatch',lambda:self.prepare())
    def test_seo_offer_changes_rejected(self):
        self.request['inputs']['approved_seo']['rows']['a']['price']='0'
        self.assertCode('scope_violation',lambda:self.prepare())
    def test_missing_approved_copy_hands_to_existing_skill(self):
        del self.request['inputs']['approved_seo']
        self.assertEqual(self.prepare()['status'],'awaiting_input')
        self.assertEqual(self.plan['body']['handoff']['skill'],'amazon-seo')
    def test_unapproved_copy_rejected(self):
        self.request['inputs']['approved_seo']['approved']=False
        self.assertCode('unapproved_copy',lambda:self.prepare())
    def test_changed_artifact_blocks_execute(self):
        result=self.prepare();Path(self.plan['body']['upload']).write_text('changed')
        self.assertCode('artifact_changed',lambda:self.service.execute(self.execute_request(result)))
    def test_wrong_grant_rejected(self):
        result=self.prepare();request=self.execute_request(result);request['grant']['targets']=['a']
        self.assertCode('grant_mismatch',lambda:self.service.execute(request))
    def test_live_execution_requires_explicit_canary(self):
        result=self.prepare();request=self.execute_request(result);request['grant'].pop('allow_live_canary')
        self.assertCode('canary_required',lambda:self.service.execute(request))
    def test_failed_executor_never_becomes_success(self):
        result=self.prepare()
        response=type('R',(),{'stdout':json.dumps({'plan_hash':result['plan_hash'],'status':'failed'}),'returncode':0})()
        with patch.object(op.subprocess,'run',return_value=response):
            self.assertEqual(self.service.execute(self.execute_request(result))['status'],'failed')
    def test_unknown_submit_result_is_not_retried(self):
        result=self.prepare()
        with patch.object(op.subprocess,'run',side_effect=op.subprocess.TimeoutExpired('node',1)) as runner:
            self.assertEqual(self.service.execute(self.execute_request(result))['status'],'uncertain')
            self.assertEqual(self.service.execute(self.execute_request(result))['status'],'uncertain')
            self.assertEqual(runner.call_count,1)
    def test_reconcile_requires_live_values(self):
        result=self.prepare();request=self.evidence_request(result)
        request['evidence']['rows']['b']['bullet_point.0.value']='Old'
        self.assertEqual(self.service.reconcile(request)['status'],'partial')
        request['evidence']['rows']=self.plan['body']['expected_rows']
        self.assertEqual(self.service.reconcile(request)['status'],'verified')
    def test_processing_is_not_completion(self):
        result=self.prepare();request=self.evidence_request(result);request['evidence']['processing_status']='processing'
        self.assertEqual(self.service.reconcile(request)['status'],'processing')
    def test_reconciliation_checks_identity_and_timestamp(self):
        result=self.prepare();request=self.evidence_request(result)
        request['evidence']['observed_at']='2000-01-01T00:00:00+00:00'
        self.assertCode('stale_evidence',lambda:self.service.reconcile(request))
        request['evidence']['observed_at']=op.now();request['evidence']['account']=dict(self.account,marketplace='DE')
        self.assertCode('identity_mismatch',lambda:self.service.reconcile(request))
    def test_no_change_is_verified_without_execution(self):
        self.request['inputs']['approved_seo']['rows']={'a':{'item_name.0.value':'Old A'},'b':{'bullet_point.0.value':'Keep B'}}
        self.assertEqual(self.prepare()['status'],'verified')
    def test_shipping_missing_inputs_does_not_guess(self):
        self.request['operation']='shipment.create';self.request['inputs']={}
        result=self.prepare()
        self.assertEqual(result['status'],'awaiting_input');self.assertIn('client_limits',result['required_inputs'])
    def shipment(self):
        self.request['operation']='shipment.create';self.request['targets']=['a']
        self.request['inputs']={'shipment_reference':'purchase-order-1','ship_from':{'address_line1':'1 Example','city':'Example','postal_code':'12345','country':'US'},'lines':[{'sku':'a','quantity':4}], 'cartons':[{'id':'box1','weight':2,'weight_unit':'kg','dimensions':[10,20,30],'dimension_unit':'cm','contents':{'a':4}}], 'carrier':'Example Carrier','currency':'USD','estimated_cost':10,'client_limits':{'account':self.account,'max_units':20,'max_cost':30,'currency':'USD','carriers':['Example Carrier']},'existing_shipments':{'account':self.account,'observed_at':op.now(),'shipments':[]},'ship_date':'2026-09-07','ship_mode':'SPD','carrier_mode':'non_partnered','packing_templates':{'a':{'name':'box','units_per_box':4}},'label_format':'thermal_4x6'}
    def test_shipment_carton_coverage(self):
        self.shipment();self.request['inputs']['cartons'][0]['contents']['a']=3
        self.assertCode('carton_coverage',lambda:self.prepare())
    def test_shipment_duplicate_reference(self):
        self.shipment();self.request['inputs']['existing_shipments']['shipments']=[{'shipment_reference':'purchase-order-1','status':'working'}]
        self.assertCode('duplicate_shipment',lambda:self.prepare())
    def test_shipment_cost_and_carrier_limits(self):
        self.shipment();self.request['inputs']['estimated_cost']=31
        self.assertCode('cost_limit',lambda:self.prepare())
        self.request['inputs']['estimated_cost']=10;self.request['inputs']['carrier']='Other'
        self.assertCode('carrier_limit',lambda:self.prepare())
    def test_shipment_artifact_label_coverage(self):
        self.shipment();result=self.prepare();request=self.evidence_request(result)
        request['evidence'].update(shipment_reference='purchase-order-1',carrier='Example Carrier',currency='USD',actual_cost=10,quantities={'a':4},labels=[])
        self.assertCode('label_coverage',lambda:self.service.reconcile(request))
    def catalog(self, operation='create_parent'):
        import openpyxl
        template=self.root/'template.xlsm';clr=self.root/'clr.xlsx'
        book=openpyxl.Workbook();sheet=book.active;sheet.title='Template'
        sheet['A1']='settings=labelRow=4;attributeRow=5;dataRow=8;contributorId=amzn1.cr.o.TESTMERCHANT'
        attrs=['item_sku','product_type','::record_action','parentage_level#1.value','child_parent_sku_relationship#1.parent_sku','variation_theme#1.name','package_size_name#1.value','item_name#1.value','brand#1.value','product_description#1.value']
        for i,field in enumerate(attrs,1):sheet.cell(4,i,field);sheet.cell(5,i,field)
        book.save(template);book.close()
        book=openpyxl.Workbook();sheet=book.active;sheet.append(['SKU','ASIN','Parent SKU']);sheet.append(['child','B000000001','parent']);sheet.append(['parent' if operation != 'create_parent' else 'other-parent','B000000002','']);book.save(clr);book.close()
        self.account.update(seller_id='TESTMERCHANT',seller_central_name='Example Seller')
        manifest={'schema_version':1,'client':'test-brand','seller_account':'Example Seller','marketplace':'US','date':'2026-09-06','operation':operation,'source':{},'family':{'variation_theme':'PACKAGE_SIZE_NAME','variation_attribute':'package_size_name#1.value','parent':{'sku':'parent','product_type':'SKIN_SERUM','status':'new','title':'Example','fields':{'brand#1.value':'Example','product_description#1.value':'Example'}},'children':[{'sku':'child','asin':'B000000001','status':'existing','variation_value':'One'}]}}
        self.request.update(operation='catalog.family.update',targets=['parent','child'],inputs={'manifest':manifest,'source_artifacts':{'blank_template':{'path':str(template),'sha256':op.file_hash(template)},'category_listings_report':{'path':str(clr),'sha256':op.file_hash(clr)}},'existing_family':{'account':self.account,'observed_at':op.now(),'relationships':{'child':'parent'}}})
    def test_catalog_builds_real_narrow_workbook(self):
        self.catalog();result=self.prepare()
        self.assertEqual(result['status'],'prepared');self.assertEqual(self.plan['body']['adapter'],'catalog.cdp')
        self.assertEqual(self.plan['body']['stages'][0]['skus'],['parent','child'])
        self.assertEqual(self.plan['body']['protected_children'],['child'])
    def test_catalog_template_account_binding(self):
        self.catalog();self.request['account']['seller_id']='OTHER'
        self.assertCode('template_owner_mismatch',lambda:self.prepare())
    def test_parent_deletion_rejects_hidden_children(self):
        self.catalog('rebuild_family');self.request['inputs']['existing_family']['relationships']['forgotten-child']='parent'
        self.assertCode('undeclared_children',lambda:self.prepare())
    def test_catalog_rebuild_waits_for_verified_first_stage(self):
        self.catalog('rebuild_family');result=self.prepare();request=self.evidence_request(result)
        request['evidence'].update(stage=2,relationships={'child':'parent'},child_offers={'child':'preserved'})
        self.assertCode('stage_order',lambda:self.service.reconcile(request))
        request['evidence'].update(stage=1,relationships={'child':''},deleted_parents=['parent'])
        first=self.service.reconcile(request);self.assertEqual(first['status'],'partial');self.assertEqual(first['next_stage'],2)
    def test_live_collector_missing_asin_mapping_preserves_pending(self):
        result=self.prepare()
        response=type('R',(),{'stdout':json.dumps({'plan_hash':result['plan_hash'],'status':'processing','submission_id':'feed-1'}),'returncode':0})()
        with patch.object(op.subprocess,'run',return_value=response):self.service.execute(self.execute_request(result))
        with patch.object(self.service, 'collect_export', return_value=None):
            result=self.service.reconcile({'schema_version':1,'operation_id':'run-1','plan_hash':result['plan_hash']})
        self.assertEqual(result['status'],'processing');self.assertFalse(result['verified'])
    def test_standalone_product_full_update_reopens_all_fields(self):
        import openpyxl
        self.catalog()
        template=Path(self.request['inputs']['source_artifacts']['blank_template']['path'])
        book=openpyxl.load_workbook(template);sheet=book['Template']
        for i,field in enumerate(['externally_assigned_product_identifier#1.value','externally_assigned_product_identifier#1.type'],11):sheet.cell(4,i,field);sheet.cell(5,i,field)
        book.save(template);book.close()
        self.request['inputs']['source_artifacts']['blank_template']['sha256']=op.file_hash(template)
        self.request['operation']='catalog.products.create';self.request['targets']=['new-product']
        manifest=self.request['inputs']['manifest'];manifest['operation']='create_products';manifest.pop('family')
        manifest['products']=[{'sku':'new-product','product_type':'SKIN_SERUM','title':'Example Product','fields':{'brand#1.value':'Example','product_description#1.value':'Product description'},'gtin_exempt':True}]
        result=self.prepare();self.assertEqual(result['status'],'prepared');self.assertEqual(self.plan['body']['stages'][0]['skus'],['new-product'])
    def test_images_create_real_scoped_flatfile_upload(self):
        from PIL import Image
        image=self.root/'image.png';Image.new('RGB',(500,500),'white').save(image)
        self.source.write_text('sku,main_product_image_locator.0.media_location\na,https://example.com/old.jpg\n')
        self.request.update(operation='listing.images.replace',targets=['a'],inputs={'images':[{'sku':'a','slot':'MAIN','artifact':{'path':str(image),'sha256':op.file_hash(image)},'url':'https://example.com/new.jpg'}],'image_fields':{'MAIN':'main_product_image_locator.0.media_location'},'live':{'account':self.account,'observed_at':op.now(),'rows':{'a':{'main_product_image_locator.0.media_location':'https://example.com/old.jpg'}}},'source_export':{'path':str(self.source),'sha256':op.file_hash(self.source)}})
        with patch.object(op, 'verify_hosted_asset', return_value={'sha256':op.file_hash(image),'match_method':'sha256','verified_at':op.now()}):
            result=self.prepare()
        self.assertEqual(result['status'],'prepared');self.assertEqual(self.plan['body']['adapter'],'flatfilepro.cdp')
    def test_carried_values_require_fresh_live_match(self):
        self.request['inputs']['live']['rows']['a']['bullet_point.0.value']='Changed elsewhere'
        self.assertCode('stale_source',lambda:self.prepare())
    def health(self):
        self.request['operation']='account_health.fields.restore'
        self.request['targets']=['a']
        self.source.write_text('sku,item_name.0.value\na,\n')
        self.request['inputs']={'finding_id':'synthetic-required-field','required_missing':{'a':['item_name.0.value']},'live':{'account':self.account,'observed_at':op.now(),'rows':{'a':{'item_name.0.value':''}}},'source_export':{'path':str(self.source),'sha256':op.file_hash(self.source)},'approved_product_data':{'account':self.account,'approved':True,'verified':True,'approved_by':'operator','source_id':'approved-product-label','verified_at':op.now(),'rows':{'a':{'item_name.0.value':'Approved title'}}}}
    def health_execute_request(self, result):
        request=self.execute_request(result);request['grant'].update(operation='account_health.fields.restore',targets=['a']);return request
    def test_account_health_requires_approved_facts_and_empty_fields(self):
        self.health();self.request['inputs']['approved_product_data']['verified']=False
        self.assertCode('unapproved_product_data',lambda:self.prepare())
        self.request['inputs']['approved_product_data']['verified']=True;self.request['inputs']['live']['rows']['a']['item_name.0.value']='Different title'
        self.assertCode('health_conflict',lambda:self.prepare())
    def test_account_health_cannot_restore_price(self):
        self.health();self.request['inputs']['required_missing']={'a':['standard_price']}
        self.assertCode('health_scope_violation',lambda:self.prepare())
    def test_account_health_preflight_report_has_no_effects_and_requeues(self):
        self.health();result=self.prepare()
        with patch.object(self.service,'collect_export',return_value=None),patch.object(op.subprocess,'run') as browser:
            state=self.service.execute(self.health_execute_request(result))
            self.assertEqual(state['status'],'processing');self.assertEqual(state['phase'],'preflight');self.assertFalse(state['effects_started']);browser.assert_not_called()
        fresh={'path':str(self.source),'sha256':op.file_hash(self.source),'report_generated_at':op.now(),'observed_at':op.now(),'rows':{'a':{'item_name.0.value':''}}}
        with patch.object(self.service,'collect_export',return_value=fresh):
            state=self.service.reconcile({'schema_version':1,'operation_id':'run-1','plan_hash':result['plan_hash']})
            self.assertEqual(state['status'],'partial');self.assertFalse(state['effects_started']);self.assertEqual(state['next_action'],'execute')
        response=type('R',(),{'stdout':json.dumps({'plan_hash':result['plan_hash'],'status':'processing','submission_id':'feed-1'}),'returncode':0})()
        with patch.object(op.subprocess,'run',return_value=response) as browser:
            state=self.service.execute(self.health_execute_request(result));self.assertEqual(state['status'],'processing');self.assertTrue(state['effects_started']);browser.assert_called_once()
    def test_account_health_conflict_during_fresh_preflight_stops(self):
        self.health();result=self.prepare()
        with patch.object(self.service,'collect_export',return_value={'rows':{'a':{'item_name.0.value':'New conflicting title'}}}):
            self.assertCode('health_conflict',lambda:self.service.execute(self.health_execute_request(result)))
    def test_account_health_carried_grid_cells_must_match_fresh_report(self):
        import openpyxl
        self.health()
        title, bullet = 'item_name.0.value', 'bullet_point.0.value'
        self.request['targets'] = ['a', 'b']
        self.source.write_text('sku,item_name.0.value,bullet_point.0.value\na,,Keep A\nb,Keep B,\n')
        inputs = self.request['inputs']
        inputs['required_missing'] = {'a': [title], 'b': [bullet]}
        inputs['live']['rows'] = {'a': {title: '', bullet: 'Keep A'}, 'b': {title: 'Keep B', bullet: ''}}
        inputs['approved_product_data']['rows'] = {'a': {title: 'Approved A'}, 'b': {bullet: 'Approved B'}}
        inputs['source_export']['sha256'] = op.file_hash(self.source)
        result = self.prepare()
        book = openpyxl.load_workbook(self.plan['body']['upload'], read_only=True)
        try:
            rows = list(book.active.values)
        finally:
            book.close()
        actual = {row[0]: dict(zip(rows[0][1:], row[1:])) for row in rows[1:]}
        self.assertEqual(actual, {'a': {title: 'Approved A', bullet: 'Keep A'}, 'b': {title: 'Keep B', bullet: 'Approved B'}})
        request = self.health_execute_request(result)
        request['grant']['targets'] = ['a', 'b']
        for sku, field in [('a', bullet), ('b', title)]:
            for value in [None, 'Independently edited value']:
                with self.subTest(sku=sku, field=field, value=value):
                    fresh_rows = copy.deepcopy(inputs['live']['rows'])
                    if value is None:
                        del fresh_rows[sku][field]
                    else:
                        fresh_rows[sku][field] = value
                    with patch.object(self.service, 'collect_export', return_value={'rows': fresh_rows}), patch.object(op.subprocess, 'run') as adapter:
                        self.assertCode('health_conflict', lambda: self.service.execute(request))
                        adapter.assert_not_called()
        response = type('R', (), {'stdout': json.dumps({'plan_hash': result['plan_hash'], 'status': 'processing', 'submission_id': 'feed-1'}), 'returncode': 0})()
        with patch.object(self.service, 'collect_export', return_value={'rows': inputs['live']['rows']}), patch.object(op.subprocess, 'run', return_value=response) as adapter:
            self.assertEqual(self.service.execute(request)['status'], 'processing')
            adapter.assert_called_once()
    def test_validated_adapter_grant_allows_normal_execution(self):
        result=self.prepare();request=self.execute_request(result);request['grant'].pop('allow_live_canary');request['grant']['allow_validated_adapter']=True
        response=type('R',(),{'stdout':json.dumps({'plan_hash':result['plan_hash'],'status':'processing','submission_id':'feed-1'}),'returncode':0})()
        with patch.object(op.subprocess,'run',return_value=response):self.assertEqual(self.service.execute(request)['status'],'processing')
    def test_hidden_field_collection_uses_fresh_backend_report(self):
        self.request['operation']='flatfile.apply';self.request['inputs'].pop('approved_seo');self.request['inputs']['desired_rows']={'a':{'item_name.0.value':'New A'},'b':{'bullet_point.0.value':'New B'}}
        result=self.prepare();state=self.service.view(self.service.directory('run-1'));state['submission_id']='feed-1'
        fresh={'source_id':'fresh-category-report','rows':self.plan['body']['expected_rows']}
        with patch.object(self.service,'collect_export',return_value=fresh) as collector:
            evidence=self.service.collect(self.service.directory('run-1'),state,self.plan)
            self.assertEqual(evidence['processing_status'],'live_observed');self.assertEqual(evidence['rows'],fresh['rows']);collector.assert_called_once()
    def test_account_health_already_restored_has_bound_json_receipt(self):
        self.health();result=self.prepare()
        fresh={'path':str(self.source),'sha256':op.file_hash(self.source),'report_generated_at':op.now(),'observed_at':op.now(),'rows':{'a':{'item_name.0.value':'Approved title'}}}
        with patch.object(self.service,'collect_export',return_value=fresh):state=self.service.execute(self.health_execute_request(result))
        self.assertTrue(state['verified']);self.assertTrue(state['no_changes']);self.assertFalse(state['effects_started'])
        receipt=json.loads(Path(state['evidence_path']).read_text())
        self.assertEqual(receipt['plan_hash'],result['plan_hash']);self.assertEqual(receipt['account'],self.account);self.assertEqual(receipt['report_artifact']['sha256'],op.file_hash(self.source))
    def test_same_image_url_requires_live_content_verification(self):
        from PIL import Image
        image=self.root/'image.png';Image.new('RGB',(500,500),'white').save(image)
        field='main_product_image_locator.0.media_location';url='https://example.com/approved.png'
        self.request.update(operation='listing.images.replace',targets=['a'],inputs={'images':[{'sku':'a','slot':'MAIN','artifact':{'path':str(image),'sha256':op.file_hash(image)},'url':url}],'image_fields':{'MAIN':field},'live':{'account':self.account,'observed_at':op.now(),'rows':{'a':{field:url}}}})
        with patch.object(op,'verify_hosted_asset',return_value={'sha256':op.file_hash(image),'match_method':'sha256','verified_at':op.now()}):result=self.prepare()
        self.assertFalse(result['verified']);self.assertTrue(self.plan['body']['no_changes'])
        execute=self.execute_request(result);execute['grant'].update(operation='listing.images.replace',targets=['a'])
        state=self.service.execute(execute);self.assertEqual(state['phase'],'verification');self.assertFalse(state['effects_started'])
        evidence={'account':self.account,'plan_hash':result['plan_hash'],'source_id':'synthetic-live-image','observed_at':op.now(),'processing_status':'live_observed','images':[{'sku':'a','slot':'MAIN','source_sha256':op.file_hash(image),'visually_verified':True,'live_url':url}]}
        state=self.service.reconcile({'schema_version':1,'operation_id':'run-1','plan_hash':result['plan_hash'],'evidence':evidence})
        self.assertTrue(state['verified'])
    def test_invalid_id_rejected(self):
        self.request['operation_id']='../escape';self.assertCode('invalid_id',lambda:self.prepare())
    def test_prepared_plan_tampering_rejected(self):
        result=self.prepare();self.plan['targets']=['outside'];Path(result['plan_path']).write_text(json.dumps(self.plan))
        self.assertCode('plan_mismatch',lambda:self.service.execute(self.execute_request(result)))
    def test_client_ready_capabilities_are_honest(self):
        self.assertTrue(all(not c['production_ready'] for c in op.capabilities()['operations'].values()))

if __name__=='__main__':unittest.main()
