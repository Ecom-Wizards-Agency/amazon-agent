#!/usr/bin/env python3
"""Amazon operation plans and restart-safe journals. JSON protocol version 1.

This local trusted-process boundary is not an authentication service. The caller
must validate the principal before supplying a grant or collected evidence.
"""
from __future__ import annotations

import argparse
import contextlib
import contextvars
import csv
import datetime as dt
import fcntl
import hashlib
import importlib.util
import io
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from urllib.parse import urlsplit

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools/browserctl"))
from browser_session import session_environment
ALIASES = {'seo.apply': 'seo.update', 'flatfile.apply': 'flatfilepro.update', 'listing.images.replace': 'listing.images', 'catalog.products.create': 'catalog.change', 'catalog.family.update': 'catalog.change', 'account_health.fields.restore': 'flatfilepro.update', 'account_health.images.restore': 'listing.images'}
OPERATIONS = {'seo.update', 'flatfilepro.update', 'listing.images', 'catalog.change', 'shipment.create', 'case.create', 'case.reply'} | set(ALIASES)

def operation_kind(value):
    return ALIASES.get(value, value)
TERMINAL = {'verified', 'failed', 'stalled'}
_CLOSE_RECONCILE_TABS = contextvars.ContextVar('close_reconcile_tabs', default=False)
IMAGE_FULL_VERIFY_SECONDS = 3600
SEO_FIELDS = re.compile(r'^(item_name|bullet_point|product_description|generic_keyword|title_differentiation)([._]|$)|^itemName$')

class OperationError(ValueError):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


def require(condition, code, message):
    if not condition:
        raise OperationError(code, message)


def canonical(value):
    return json.dumps(value, sort_keys=True, separators=(',', ':'), ensure_ascii=False, allow_nan=False).encode()


def digest(value):
    return hashlib.sha256(canonical(value)).hexdigest()


def file_hash(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def now():
    return dt.datetime.now(dt.timezone.utc).isoformat()


def timestamp(value):
    try:
        result = dt.datetime.fromisoformat(value.replace('Z', '+00:00'))
        require(result.tzinfo is not None, 'invalid_time', 'Timestamps require a timezone')
        return result
    except (ValueError, TypeError, AttributeError) as exc:
        raise OperationError('invalid_time', 'Expected an ISO timestamp with timezone') from exc


def atomic_json(path, value):
    path = Path(path)
    fd, name = tempfile.mkstemp(prefix='.write-', dir=path.parent)
    try:
        with os.fdopen(fd, 'wb') as stream:
            stream.write(canonical(value) + b'\n')
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def load_module(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def evidence_tools():
    return load_module('merlin_operation_evidence', 'tools/amazon-operations/evidence.py')


def case_tools():
    return load_module('amazon_case_operations', 'tools/amazon-operations/case_operations.py')


def verify_hosted_asset(path, url):
    module = evidence_tools()
    data = module.fetch_public_image(url)
    method = module.same_image_content(Path(path).read_bytes(), data)
    require(method is not None, 'hosted_image_mismatch', 'Hosted image does not match the approved local image bytes or decoded pixels')
    return {'sha256': hashlib.sha256(data).hexdigest(), 'match_method': method, 'verified_at': now()}


def health_request(request):
    if not request['operation'].startswith('account_health.'):
        return request
    inputs, acct = request['inputs'], request['account']
    approved = inputs.get('approved_product_data')
    scoped(approved, acct, 'approved product data')
    require(approved.get('approved') is True and approved.get('verified') is True and approved.get('source_id') and approved.get('approved_by'), 'unapproved_product_data', 'Routine restoration requires approved verified product data and provenance')
    timestamp(approved.get('verified_at'))
    require(inputs.get('finding_id') and isinstance(inputs.get('required_missing'), dict) and set(inputs['required_missing']) == set(request['targets']), 'missing_health_scope', 'Account Health finding and exact required-missing field scope are required')
    live = inputs.get('live')
    scoped(live, acct, 'live missing-field evidence')
    desired = {}
    for sku, fields in inputs['required_missing'].items():
        require(isinstance(fields, list) and fields, 'missing_health_scope', 'Each affected SKU requires exact missing fields')
        desired[sku] = {}
        for field in fields:
            require(not re.search(r'^(price|standard_price|purchasable_offer|quantity|fulfillment|condition|parent|child_parent|variation|product_type|record_action|item_sku)', field, re.I), 'health_scope_violation', 'Routine restoration cannot change offers, inventory, identity or relationships')
            current = live.get('rows', {}).get(sku, {}).get(field, '__unknown__')
            value = approved.get('rows', {}).get(sku, {}).get(field)
            require(value is not None and as_text(value).strip(), 'missing_approved_value', 'Approved source lacks a required nonempty restoration value')
            require(current != '__unknown__', 'missing_live_field', 'Current field absence must be positively observed')
            require(not as_text(current).strip() or as_text(current) == as_text(value), 'health_conflict', 'Routine restoration cannot overwrite a nonempty conflicting field')
            desired[sku][field] = value
    if inputs.get('desired_rows') is not None:
        require(inputs['desired_rows'] == desired, 'health_source_mismatch', 'Requested restoration differs from approved product facts')
    inputs['desired_rows'] = desired
    inputs['scope_fields'] = list(dict.fromkeys(field for fields in inputs['required_missing'].values() for field in fields))
    if request['operation'] == 'account_health.images.restore':
        approved_images = {(x.get('sku'), x.get('slot')): x for x in approved.get('images', [])}
        for image in inputs.get('images', []):
            source = approved_images.get((image.get('sku'), image.get('slot')), {})
            require(source.get('url') == image.get('url') and source.get('sha256') == image.get('artifact', {}).get('sha256'), 'health_source_mismatch', 'Image restoration must use the exact approved hosted asset')
            field = inputs.get('image_fields', {}).get(image.get('slot'))
            require(desired.get(image.get('sku'), {}).get(field) == image.get('url'), 'health_source_mismatch', 'Approved image and approved field restoration values differ')
    return request


def account(value):
    require(isinstance(value, dict), 'missing_account', 'account is required')
    for key in ('client_slug', 'profile_key', 'marketplace'):
        require(isinstance(value.get(key), str) and value[key].strip(), 'missing_account', f'account.{key} is required')
    return value


def scoped(document, expected, label):
    require(isinstance(document, dict) and document.get('account') == expected,
            'identity_mismatch', f'{label} account must exactly match the operation account')


def input_artifact(spec, directory, name):
    require(isinstance(spec, dict) and spec.get('path') and spec.get('sha256'), 'missing_artifact', f'{name} needs path and sha256')
    source = Path(spec['path']).expanduser().resolve()
    require(source.is_file(), 'missing_artifact', f'{name} file is missing')
    require(file_hash(source) == spec['sha256'], 'artifact_changed', f'{name} checksum mismatch')
    target = directory / (name + source.suffix.lower())
    shutil.copyfile(source, target)
    require(file_hash(target) == spec['sha256'], 'artifact_changed', f'{name} changed while copied')
    return target


def as_text(value):
    require(value is None or isinstance(value, (str, int, float)) and not isinstance(value, bool), 'invalid_value', 'Field values must be text or numbers')
    if isinstance(value, float):
        require(math.isfinite(value), 'invalid_value', 'Field values must be finite')
    return '' if value is None else str(value)


def changed_fields(desired, live, targets, fields):
    require(isinstance(fields, list) and fields and len(set(fields)) == len(fields), 'missing_scope', 'scope_fields must be a nonempty unique list')
    require(set(desired) <= set(targets), 'scope_violation', 'Desired rows include unrequested SKUs')
    result = []
    for sku in targets:
        require(sku in desired and sku in live, 'missing_live_row', f'Both desired and current rows are required for {sku}')
        require(set(desired[sku]) <= set(fields), 'scope_violation', f'Unrequested fields for {sku}')
        for field, value in desired[sku].items():
            require(field in live[sku], 'missing_live_field', f'Current {sku}/{field} is unknown')
            before, after = as_text(live[sku][field]), as_text(value)
            if before != after:
                result.append({'sku': sku, 'field': field, 'before': before, 'after': after})
    return result


def prepare_copy(request, directory):
    inputs, acct = request['inputs'], request['account']
    live = inputs.get('live')
    scoped(live, acct, 'live listing evidence')
    timestamp(live.get('observed_at'))
    require(timestamp(live['observed_at']) <= timestamp(now()) + dt.timedelta(minutes=5), 'invalid_time', 'Live evidence cannot be future dated')
    if operation_kind(request['operation']) == 'seo.update':
        approved = inputs.get('approved_seo')
        if not approved:
            return {'status': 'awaiting_input', 'required_inputs': ['approved_seo'], 'handoff': {'skill': 'amazon-seo', 'reason': 'Run existing product-facts, research and SEO workflow; return reviewed copy and source provenance'}}
        scoped(approved, acct, 'approved SEO')
        require(approved.get('approved') is True and approved.get('source_id') and approved.get('approved_by'), 'unapproved_copy', 'SEO requires approved copy with source_id and approved_by')
        timestamp(approved.get('approved_at'))
        require(approved.get('superseded') is not True, 'superseded_copy', 'Do not apply superseded SEO')
        desired = approved.get('rows', {})
        require(all(SEO_FIELDS.search(field) for row in desired.values() for field in row), 'scope_violation', 'SEO cannot modify offer, inventory, image or relationship fields')
    else:
        desired = inputs.get('desired_rows', {})
    changes = changed_fields(desired, live.get('rows', {}), request['targets'], inputs.get('scope_fields'))
    if not changes:
        return {'status': 'verified', 'no_changes': True, 'changes': [], 'expected_rows': desired, 'verification': 'Current supplied listing evidence already matches requested fields', 'observed_at': live['observed_at'], 'source_id': live.get('source_id', 'supplied-live-listing-evidence')}
    source = input_artifact(inputs.get('source_export'), directory, 'source-export')
    prep = load_module('merlin_ffp_prep', 'skills/amazon-flatfilepro/scripts/prepare_flatfilepro_upload.py')
    headers = prep.read_source_headers(source)
    source_rows = prep.read_source_rows(source, headers)
    require(len(headers) == len(set(headers)), 'ambiguous_headers', 'Source export has duplicate headers')
    canon = [prep.canonical_attribute(h) for h in headers]
    require(len(canon) == len(set(canon)), 'ambiguous_headers', 'Source headers canonicalize to duplicate attributes')
    for item in changes:
        require(item['field'] in headers, 'unknown_field', f"Exact source header unavailable: {item['field']}")
        require(item['sku'] in source_rows, 'missing_source_row', f"Source export missing SKU: {item['sku']}")
        require(as_text(source_rows[item['sku']].get(item['field'])) == item['before'], 'stale_source', 'Source export does not match current listing evidence')
    transmitted = list(dict.fromkeys(c['field'] for c in changes))
    companions = {'unit_count.0.value': ['unit_count.0.type.value'], 'unit_count.0.type.value': ['unit_count.0.value']}
    for prefix in ('item_weight', 'item_package_weight', 'item_display_weight'):
        companions[prefix + '.0.value'] = [prefix + '.0.unit']
        companions[prefix + '.0.unit'] = [prefix + '.0.value']
    for field in list(transmitted):
        for required in companions.get(field, []):
            require(required in headers and required in inputs['scope_fields'], 'missing_companion', f'Exact scoped companion field required: {required}')
            if required not in transmitted:
                transmitted.append(required)
    expected_skus = {c['sku'] for c in changes}
    for sku in expected_skus:
        for field in transmitted:
            require(field in live['rows'][sku] and as_text(source_rows[sku].get(field)) == as_text(live['rows'][sku][field]), 'stale_source', f'Carried field {sku}/{field} does not match live evidence')
    changes_path = directory / 'changes.csv'
    with changes_path.open('w', newline='', encoding='utf-8') as stream:
        writer = csv.DictWriter(stream, fieldnames=['sku', 'attribute', 'value'])
        writer.writeheader()
        writer.writerows({'sku': c['sku'], 'attribute': c['field'], 'value': c['after']} for c in changes)
        first_sku = changes[0]['sku']
        writer.writerows({'sku': first_sku, 'attribute': field, 'value': source_rows[first_sku][field]} for field in transmitted if field not in {c['field'] for c in changes})
    output = directory / 'upload.xlsx'
    with contextlib.redirect_stdout(io.StringIO()):
        result = prep.main(['--source', str(source), '--changes', str(changes_path), '--output', str(output), '--fill-unchanged'])
    require(result == 0, 'preparation_failed', 'FlatFilePro preparation failed')
    preserve_headers = inputs.get('image_policy') == 'secondary_slots_only'
    if preserve_headers:
        # The current FFP image selector exposes the export's __1__ spelling.
        # Preserve this narrow batch's exact template headers; legacy copy
        # workbooks retain their established canonical-header behavior.
        from openpyxl import load_workbook
        raw_headers = {prep.canonical_attribute(field): field for field in transmitted}
        workbook = load_workbook(output)
        try:
            for cell in workbook.active[1]:
                if cell.value != 'sku':
                    require(cell.value in raw_headers, 'artifact_mismatch', 'Unexpected image workbook header')
                    cell.value = raw_headers[cell.value]
            workbook.save(output)
        finally:
            workbook.close()
    upload_headers = prep.read_source_headers(output)
    upload_rows = prep.read_source_rows(output, upload_headers)
    expected_skus = {c['sku'] for c in changes}
    require(set(upload_rows) == expected_skus, 'artifact_mismatch', 'Upload row coverage differs from changed SKUs')
    output_field = (lambda field: field) if preserve_headers else prep.canonical_attribute
    expected = {sku: {output_field(field): as_text(source_rows[sku].get(field)) for field in transmitted} for sku in expected_skus}
    for c in changes:
        expected[c['sku']][output_field(c['field'])] = c['after']
    for sku, row in expected.items():
        require({k: v for k, v in upload_rows[sku].items() if k != 'sku'} == row, 'artifact_mismatch', 'Reopened upload does not match full-grid expected values')
    return {'status': 'prepared', 'changes': changes, 'expected_rows': expected, 'upload': str(output), 'mapping': [h for h in upload_headers if h != 'sku'], 'adapter': 'flatfilepro.cdp', 'requires_live_canary': True, 'sku_asins': inputs.get('sku_asins', {}), 'restore_missing_only': request['operation'].startswith('account_health.'), 'finding_id': inputs.get('finding_id')}


def prepare_images(request, directory):
    inputs = request['inputs']
    manifest = inputs.get('images')
    require(isinstance(manifest, list) and manifest, 'missing_images', 'images manifest is required')
    changes, seen = [], set()
    for image in manifest:
        sku, slot = image.get('sku'), image.get('slot')
        require(sku in request['targets'] and isinstance(slot, str) and re.fullmatch(r'MAIN|PT0[1-9]|SWCH', slot), 'scope_violation', 'Image needs a requested SKU and MAIN/PT01..PT09/SWCH slot')
        require((sku, slot) not in seen, 'duplicate_image', 'Duplicate SKU/image slot')
        seen.add((sku, slot))
        path = input_artifact(image.get('artifact'), directory, f'image-{len(changes)}')
        from PIL import Image
        with Image.open(path) as im:
            require(im.format in {'JPEG', 'PNG'} and min(im.size) >= 500, 'invalid_image', 'Images must be JPEG/PNG with both dimensions at least 500px')
            dimensions = list(im.size)
            im.verify()
        url = image.get('url', '')
        parsed = urlsplit(url)
        require(parsed.scheme == 'https' and parsed.hostname and not parsed.username and not parsed.password, 'invalid_image_url', 'Image upload URLs must be HTTPS without embedded credentials')
        hosted = verify_hosted_asset(path, url)
        changes.append({'hosted_verification': hosted, 'sku': sku, 'slot': slot, 'sha256': file_hash(path), 'path': str(path), 'url': url, 'dimensions': dimensions})
    require({c['sku'] for c in changes} == set(request['targets']), 'missing_images', 'Every target requires an image')
    fields = inputs.get('image_fields', {})
    require(all(image['slot'] in fields for image in changes), 'missing_image_mapping', 'image_fields must map each slot to an exact source-export header')
    canonical_fields = [evidence_tools().canonical_header(field) for field in fields.values()]
    require(len(set(canonical_fields)) == len(fields), 'duplicate_image_mapping', 'Each image slot must map to its own template attribute')
    for slot, field in fields.items():
        require(isinstance(field, str) and re.fullmatch(r'MAIN|PT0[1-9]|SWCH', slot), 'scope_violation', 'Unsupported image slot or attribute')
        prefix = {'MAIN': 'main_product_image_locator', 'SWCH': 'swatch_product_image_locator'}.get(slot, f'other_product_image_locator_{int(slot[2:])}' if slot.startswith('PT') else '')
        require(evidence_tools().canonical_header(field) == prefix + '.0.media_location', 'scope_violation', 'Image slot must map to its exact supported template media_location attribute')
    policy = inputs.get('image_policy')
    require(policy in {None, 'secondary_slots_only'}, 'scope_violation', 'Unknown image policy')
    if policy:
        require(all(image['slot'].startswith('PT') for image in changes) and all(slot.startswith('PT') for slot in fields), 'scope_violation', 'Secondary image updates cannot change MAIN or swatches')
    copy_request = json.loads(json.dumps(request))
    copy_request['operation'] = 'flatfilepro.update'
    # FFP can export mangled headers while the current catalog uses canonical ones.
    for row in copy_request['inputs'].get('live', {}).get('rows', {}).values():
        for field in fields.values():
            canonical_field = evidence_tools().canonical_header(field)
            if field not in row and canonical_field in row:
                row[field] = row[canonical_field]
    desired = {}
    for image in changes:
        desired.setdefault(image['sku'], {})[fields[image['slot']]] = image['url']
    copy_request['inputs']['desired_rows'] = desired
    copy_request['inputs']['scope_fields'] = list(dict.fromkeys(fields[image['slot']] for image in changes))
    prepared = prepare_copy(copy_request, directory)
    if prepared['status'] == 'verified':
        prepared.update(status='prepared', adapter='image.verify', no_changes=True, requires_live_canary=False)
    prepared['images'] = changes
    prepared['restore_missing_only'] = request['operation'].startswith('account_health.')
    prepared['sku_asins'] = inputs.get('sku_asins', {})
    if policy:
        prepared['image_policy'] = policy
        prepared['image_before_rows'] = image_before_rows(request, directory)
        prepared['image_identity_rows'] = {sku: {**dict.fromkeys(IMAGE_IDENTITY_KEYS, ''), **image_listing_identity(inputs['live']['rows'][sku])} for sku in request['targets']}
        source_kind = inputs.get('image_catalog_source')
        require(source_kind in {None, 'flatfilepro_listing_read'}, 'invalid_source', 'Unknown image catalog source')
        if source_kind:
            prepared['image_catalog_source'] = source_kind
            prepared['image_preview_metadata'] = {sku: {'itemName': inputs['live']['rows'][sku].get('itemName') or inputs['live']['rows'][sku].get('item_name.0.value'), 'productType': inputs['live']['rows'][sku].get('product_type')} for sku in request['targets']}
            require(all(x['itemName'] and x['productType'] for x in prepared['image_preview_metadata'].values()), 'missing_metadata', 'FFP image preview requires observed title and product type')
    return prepared


IMAGE_IDENTITY_KEYS = [
    'sku',
    'asin',
    'mpn',
    'color',
    'color_code',
    'size',
    'parentage',
    'parent_sku',
    'listing_relationship_evidence',
    'product_type',
    'archived',
    'itemName',
    'model_name',
    'model_number',
    'item_name.0.value',
    'model_name.0.value',
    'model_number.0.value',
    'part_number.0.value',
    'color.0.value',
    'size.0.value',
    'parentage_level.0.value',
    'child_parent_sku_relationship.0.parent_sku',
    'externally_assigned_product_identifier.0.value',
    'merchant_suggested_asin.0.value',
]


def image_listing_identity(row):
    """Freeze every observed identity field, preserving blank optional values."""
    attributes = re.compile(r'^(?:item_name|model_name|model_number|part_number|color|size|'
                            r'parentage_level|child_parent_sku_relationship|'
                            r'externally_assigned_product_identifier|merchant_suggested_asin)\.')
    return {k: v for k, v in row.items() if k in IMAGE_IDENTITY_KEYS or attributes.match(k)}


def image_before_rows(request, directory):
    """Freeze all exported image attributes, including untouched slots and MAIN."""
    inputs = request['inputs']
    source = input_artifact(inputs.get('source_export'), directory, 'image-baseline-export')
    evidence = evidence_tools()
    rows = evidence.report_rows(source)
    baseline = {}
    for sku in request['targets']:
        row = rows.get(sku, {})
        live = inputs['live']['rows'].get(sku, {})
        image_fields = {field: as_text(value) for field, value in row.items()
                        if re.fullmatch(r'(?:main_product_image_locator|swatch_product_image_locator|other_product_image_locator_[1-9])\.0\.media_location', field)}
        require(image_fields.get('main_product_image_locator.0.media_location'), 'missing_main_image', 'Secondary updates require an observed existing MAIN image')
        require(all(field in live and as_text(live[field]) == value for field, value in image_fields.items()), 'stale_source', 'Complete live image evidence must match the source export, including protected slots')
        asin = inputs.get('sku_asins', {}).get(sku)
        require(isinstance(asin, str) and re.fullmatch(r'[A-Z0-9]{10}', asin) and (live.get('asin') or live.get('ASIN')) == asin, 'image_identity_mismatch', 'Live catalog evidence must bind every seller SKU to its exact ASIN')
        parentage = evidence.field_value(live, 'parentage')
        parent = evidence.field_value(live, 'parent_sku')
        standalone = (inputs.get('allow_standalone_images') is True and inputs.get('image_catalog_source') == 'flatfilepro_listing_read' and
                      not parentage and not parent and live.get('listing_relationship_evidence') == 'standalone')
        require(parentage == 'child' or standalone, 'image_identity_mismatch', 'Secondary updates require verified child listings or explicitly scoped standalone listings')
        baseline[sku] = {**image_fields, 'asin': asin, 'parentage_level.0.value': parentage,
                         'child_parent_sku_relationship.0.parent_sku': parent}
        if standalone:
            baseline[sku]['listing_relationship_evidence'] = 'standalone'
    return baseline


def prepare_catalog(request, directory):
    manifest = json.loads(json.dumps(request['inputs'].get('manifest', {})))
    acct = request['account']
    seller_account = acct.get('seller_account') or acct.get('seller_central_name')
    merchant_id = acct.get('merchant_id') or acct.get('seller_id')
    require(manifest.get('marketplace') == acct['marketplace'] and manifest.get('client') == acct['client_slug'], 'identity_mismatch', 'Catalog manifest client/marketplace mismatch')
    require(seller_account and manifest.get('seller_account') == seller_account, 'identity_mismatch', 'Catalog seller account mismatch')
    require(merchant_id, 'missing_merchant_id', 'Verified merchant_id is required for template ownership')
    sources = request['inputs'].get('source_artifacts', {})
    for key in ('category_listings_report', 'blank_template'):
        copied = input_artifact(sources.get(key), directory, key)
        manifest.setdefault('source', {})[key] = str(copied)
    import openpyxl
    book = openpyxl.load_workbook(manifest['source']['blank_template'], read_only=True, keep_vba=False)
    try:
        settings = str(book['Template']['A1'].value)
        match = re.search(r'contributorId=amzn1\.cr\.o\.([^;,&\s]+)', settings)
        require(match and match.group(1) == merchant_id, 'template_owner_mismatch', 'Template contributorId does not match verified merchant_id')
    finally:
        book.close()
    module = load_module('merlin_catalog_pack', 'tools/amazon-catalog-change-pack/catalog_change_pack.py')
    manifest['output_dir'] = str(directory / 'catalog-pack')
    manifest['_manifest_path'] = str(directory / 'catalog-manifest.json')
    if manifest.get('operation') == 'create_products':
        return prepare_standalone_products(request, directory, manifest, module)
    specs = module.build_specs(manifest)
    # Exact scope includes every affected parent and child, including indirectly detached children.
    family = manifest.get('family', {})
    affected = {str(c['sku']) for c in family.get('children', [])}
    affected.update(family.get('removed_child_skus', []))
    affected.update(filter(None, [family.get('old_parent_sku'), family.get('parent', {}).get('sku')]))
    if manifest['operation'] in {'delete_parent', 'rebuild_family', 'detach_children', 'reparent_children'}:
        current_family = request['inputs'].get('existing_family')
        scoped(current_family, acct, 'current family evidence')
        timestamp(current_family.get('observed_at'))
        old_parent = family.get('old_parent_sku') or family['parent']['sku']
        actual_children = {sku for sku, parent in current_family.get('relationships', {}).items() if parent == old_parent}
        declared_children = {c['sku'] for c in family.get('children', [])} | set(family.get('removed_child_skus', []))
        require(actual_children <= declared_children, 'undeclared_children', 'Parent deletion would detach children missing from the explicit plan')
    require(affected <= set(request['targets']), 'scope_violation', 'Catalog plan touches an unrequested parent or child')
    atomic_json(directory / 'catalog-manifest.json', manifest)
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            pack = module.build(manifest)
    except module.CatalogError as exc:
        raise OperationError('invalid_catalog', str(exc)) from exc
    files = sorted((pack / '03-upload-ready').glob('*.xlsm'))
    require(len(files) == len(specs), 'artifact_mismatch', 'Catalog stage coverage mismatch')
    report = evidence_tools().report_rows(manifest['source']['category_listings_report'])
    protected = sorted({c['sku'] for c in family.get('children', []) if c.get('status') == 'existing'} | set(family.get('removed_child_skus', [])))
    offer_baseline = {sku: {field: value for field in ('price', 'quantity', 'fulfillment_channel', 'condition') if (value := evidence_tools().field_value(report.get(sku, {}), field)) is not None} for sku in protected}
    return {'status': 'prepared', 'manifest': manifest, 'protected_offer_values': offer_baseline, 'stages': [{'index': i + 1, 'upload': str(path), 'requires_prior_verified': i > 0, 'skus': [row['sku'] for row in specs[i][1]], 'full_update_values': {row['sku']: {k: v for k, v in row.items() if k not in {'sku', 'action'} and v is not None} for row in specs[i][1] if row.get('action') == module.FULL_UPDATE}} for i, path in enumerate(files)], 'protected_children': sorted({c['sku'] for c in family.get('children', []) if c.get('status') == 'existing'} | set(family.get('removed_child_skus', []))), 'adapter': 'catalog.cdp', 'requires_live_canary': True}


def prepare_standalone_products(request, directory, manifest, module):
    """Full-update new standalone rows using the established template machinery."""
    products = manifest.get('products', [])
    require(isinstance(products, list) and products, 'missing_products', 'create_products requires products')
    skus = [p.get('sku') for p in products]
    require(len(skus) == len(set(skus)) and set(skus) == set(request['targets']), 'scope_violation', 'New products must exactly match requested unique SKUs')
    existing = module.read_report_skus(Path(manifest['source']['category_listings_report']))
    require(not set(skus) & existing, 'existing_product', 'Full updates cannot replace existing products')
    require(len({p.get('product_type') for p in products}) == 1, 'mixed_templates', 'Create separate packs for different product types')
    rows = []
    for product in products:
        require(product.get('product_type') and product.get('title') and product.get('fields'), 'missing_product_fields', 'New products require product_type, title and exact template fields')
        fields = product['fields']
        require(not any(re.search(r'parent|variation|record_action|item_sku|contribution_sku', k, re.I) for k in fields), 'scope_violation', 'Standalone fields cannot override identity, action or relationships')
        row = {**fields, 'sku': product['sku'], 'product_type': product['product_type'], 'action': module.FULL_UPDATE, 'title': product['title']}
        if product.get('gtin_exempt') is True:
            require(not product.get('product_id'), 'invalid_gtin', 'GTIN exemption requires a blank product ID')
            row.update(external_product_id=None, external_product_id_type='GTIN Exempt')
        else:
            require(product.get('product_id') and product.get('product_id_type'), 'missing_gtin', 'A product ID/type or evidenced GTIN exemption is required')
            row.update(external_product_id=product['product_id'], external_product_id_type=product['product_id_type'])
        rows.append(row)
    output = directory / '01-create-products.xlsm'
    try:
        module.write_workbook(Path(manifest['source']['blank_template']), output, rows)
        info = module.load_template(output)
        try:
            for offset, row in enumerate(rows):
                for field, value in row.items():
                    column = module.resolve_column(info, field)
                    require(info.sheet.cell(info.data_row + offset, column).value == value, 'artifact_mismatch', 'New-product workbook differs from intended full update')
            require(not any(cell.value is not None for row in info.sheet.iter_rows(min_row=info.data_row + len(rows)) for cell in row), 'artifact_mismatch', 'Unexpected extra populated template rows')
        finally:
            module.close_workbook(info.workbook)
    except module.CatalogError as exc:
        raise OperationError('invalid_catalog', str(exc)) from exc
    return {'status': 'prepared', 'manifest': manifest, 'products': products, 'stages': [{'index': 1, 'upload': str(output), 'skus': skus, 'requires_prior_verified': False, 'full_update_values': {row['sku']: {k: v for k, v in row.items() if k not in {'sku', 'action'} and v is not None} for row in rows}}], 'protected_children': [], 'adapter': 'catalog.cdp', 'requires_live_canary': True}


def positive_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and value > 0


def prepare_shipment(request, directory):
    inputs = request['inputs']
    required = ['shipment_reference', 'ship_from', 'lines', 'cartons', 'carrier', 'currency', 'estimated_cost', 'client_limits', 'existing_shipments', 'ship_date', 'ship_mode', 'carrier_mode', 'packing_templates', 'label_format']
    missing = [name for name in required if inputs.get(name) is None or inputs.get(name) == '']
    if missing:
        return {'status': 'awaiting_input', 'required_inputs': missing}
    limits = inputs['client_limits']
    scoped(limits, request['account'], 'client shipment limits')
    missing_limits = [f'client_limits.{key}' for key in ('max_units', 'max_cost', 'currency', 'carriers') if limits.get(key) is None]
    if missing_limits:
        return {'status': 'awaiting_input', 'required_inputs': missing_limits}
    lines = inputs['lines']
    require(isinstance(lines, list) and lines, 'missing_lines', 'Shipment lines are required')
    quantities = {}
    for line in lines:
        sku, qty = line.get('sku'), line.get('quantity')
        require(sku in request['targets'] and sku not in quantities, 'scope_violation', 'Shipment lines require unique requested SKUs')
        require(isinstance(qty, int) and not isinstance(qty, bool) and qty > 0, 'invalid_quantity', 'Shipment quantities must be positive integers')
        quantities[sku] = qty
    require(set(quantities) == set(request['targets']), 'missing_lines', 'Every target requires a shipment quantity')
    require(positive_number(limits['max_units']) and sum(quantities.values()) <= limits['max_units'], 'quantity_limit', 'Shipment exceeds client unit limit')
    cost = inputs['estimated_cost']
    require(isinstance(cost, (int, float)) and not isinstance(cost, bool) and math.isfinite(cost) and cost >= 0 and positive_number(limits['max_cost']), 'invalid_cost', 'Cost and saved maximum must be finite nonnegative values')
    require(cost <= limits['max_cost'] and inputs['currency'] == limits['currency'], 'cost_limit', 'Shipment exceeds saved cost or currency limit')
    require(inputs['carrier'] in limits['carriers'], 'carrier_limit', 'Carrier is outside saved client choices')
    require(isinstance(inputs['ship_from'], dict) and all(inputs['ship_from'].get(key) for key in ('address_line1', 'city', 'postal_code', 'country')), 'missing_address', 'Complete ship-from address is required')
    packed, boxes = {sku: 0 for sku in quantities}, set()
    for carton in inputs['cartons']:
        require(carton.get('id') and carton['id'] not in boxes, 'duplicate_carton', 'Carton IDs must be unique')
        boxes.add(carton['id'])
        require(positive_number(carton.get('weight')) and carton.get('weight_unit') in {'kg', 'lb'}, 'invalid_carton', 'Carton weight and unit required')
        dims = carton.get('dimensions', [])
        require(len(dims) == 3 and all(positive_number(x) for x in dims) and carton.get('dimension_unit') in {'cm', 'in'}, 'invalid_carton', 'Three positive carton dimensions and unit required')
        for sku, qty in carton.get('contents', {}).items():
            require(sku in packed and isinstance(qty, int) and not isinstance(qty, bool) and qty > 0, 'invalid_carton', 'Carton content is outside shipment scope')
            packed[sku] += qty
    require(boxes and packed == quantities, 'carton_coverage', 'Carton quantities must exactly cover shipment lines')
    existing = inputs['existing_shipments']
    scoped(existing, request['account'], 'existing shipment evidence')
    timestamp(existing.get('observed_at'))
    for shipment in existing.get('shipments', []):
        if shipment.get('shipment_reference') == inputs['shipment_reference'] and shipment.get('status') != 'cancelled':
            raise OperationError('duplicate_shipment', 'Existing shipment already uses this shipment_reference; reconcile it instead')
    return {'status': 'prepared', 'shipment': {k: v for k, v in inputs.items() if k not in {'existing_shipments', 'client_limits'}}, 'limits': limits, 'adapter': 'shipment.cdp', 'requires_live_canary': True}


class Operations:
    def __init__(self, state_dir, *, case_state_dir=None, case_policy_path=None):
        self.root = Path(state_dir).expanduser().resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        # Constructor injection is for trusted tests/embedders; request payloads
        # and the CLI cannot redirect case delivery into another claim journal.
        self.case_root = Path(case_state_dir or Path.home() / '.amazon-agent/cases').expanduser().resolve()
        self.case_policy_path = case_policy_path

    def case_journal_boundary(self):
        require(self.root == self.case_root / 'operations', 'case_journal_mismatch', 'Case operations must use the canonical shared case journal directory')

    def directory(self, operation_id):
        require(isinstance(operation_id, str) and re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]{0,99}', operation_id), 'invalid_id', 'operation_id must be a simple 1..100 character identifier')
        path = self.root / operation_id
        path.mkdir(exist_ok=True)
        require(path.resolve().parent == self.root, 'unsafe_state', 'Operation directory must not be a symlink outside state root')
        return path

    @contextlib.contextmanager
    def lock(self, directory):
        with (directory / '.lock').open('a') as stream:
            fcntl.flock(stream, fcntl.LOCK_EX)
            yield

    def view(self, directory):
        return json.loads((directory / 'journal.json').read_text())

    def persist(self, directory, state, status, **extra):
        state.update(extra, status=status, verified=status == 'verified', updated_at=now())
        state.setdefault('events', []).append({'status': status, 'at': state['updated_at']})
        atomic_json(directory / 'journal.json', state)
        return state

    def prepare(self, request):
        require(request.get('schema_version') == 1 and request.get('operation') in OPERATIONS, 'invalid_request', 'schema_version 1 and a supported operation are required')
        account(request.get('account'))
        request = json.loads(json.dumps(request))
        raw_targets = request.get('targets')
        require(isinstance(raw_targets, list), 'invalid_targets', 'targets must be a list')
        is_case = request['operation'].startswith('case.')
        if is_case:
            self.case_journal_boundary()
        targets = case_tools().validate_targets(request['operation'], raw_targets) if is_case else [x.get('sku') if isinstance(x, dict) else x for x in raw_targets]
        request['targets'] = targets
        target_asins = {x['sku']: x['asin'] for x in raw_targets if isinstance(x, dict) and x.get('sku') and x.get('asin')}
        if target_asins:
            request.setdefault('inputs', {}).setdefault('sku_asins', target_asins)
        if not is_case:
            require(isinstance(targets, list) and targets and all(isinstance(x, str) and x for x in targets) and len(targets) == len(set(targets)), 'invalid_targets', 'targets must be unique nonempty SKU strings')
        require(isinstance(request.get('inputs'), dict), 'missing_inputs', 'inputs object required')
        request = health_request(request)
        directory = self.directory(request.get('operation_id'))
        with self.lock(directory):
            fingerprint = digest({k: v for k, v in request.items() if k not in {'task_key', 'complete_task'}})
            if (directory / 'plan.json').exists():
                old = json.loads((directory / 'plan.json').read_text())
                require(old['request_hash'] == fingerprint, 'immutable_request', 'An operation ID cannot be reused with different inputs; create a new revision ID')
                state = self.view(directory)
                if request.get('task_key'):
                    state['task_key'] = request['task_key']
                    atomic_json(directory / 'journal.json', state)
                return state
            handlers = {'seo.update': prepare_copy, 'flatfilepro.update': prepare_copy, 'listing.images': prepare_images, 'catalog.change': prepare_catalog, 'shipment.create': prepare_shipment, 'case.create': case_tools().prepare, 'case.reply': case_tools().prepare}
            body = handlers[operation_kind(request['operation'])](request, directory)
            artifacts = [{'path': str(p), 'sha256': file_hash(p)} for p in sorted(directory.rglob('*')) if p.is_file() and p.name != '.lock']
            plan = {'schema_version': 1, 'operation_id': request['operation_id'], 'operation': request['operation'], 'account': request['account'], 'targets': targets, 'request_hash': fingerprint, 'prepared_at': now(), 'body': body, 'artifacts': artifacts}
            atomic_json(directory / 'plan.json', plan)
            initial = {'schema_version': 1, 'operation_id': request['operation_id'], 'operation': request['operation'], 'plan_hash': digest(plan), 'plan_path': str(directory / 'plan.json'), 'account': request['account'], 'targets': targets, 'required_inputs': body.get('required_inputs', []), 'capability': {'adapter': body.get('adapter'), 'requires_live_canary': body.get('requires_live_canary', False)}}
            if request.get('task_key'):
                initial['task_key'] = request['task_key']
            if body['status'] == 'verified':
                evidence = {'schema_version': 1, 'account': plan['account'], 'plan_hash': digest(plan), 'source_id': body['source_id'], 'observed_at': body['observed_at'], 'processing_status': 'live_observed', 'no_changes': True, 'rows': body['expected_rows']}
                path = directory / f'evidence-{digest(evidence)}.json'
                atomic_json(path, evidence)
                initial.update(no_changes=True, effects_started=False, evidence_path=str(path))
            return self.persist(directory, initial, body['status'])

    def bound(self, directory, request):
        state = self.view(directory)
        plan = json.loads((directory / 'plan.json').read_text())
        require(request.get('schema_version') == 1 and request.get('plan_hash') == state['plan_hash'] == digest(plan), 'plan_mismatch', 'Plan hash mismatch')
        for artifact in plan['artifacts']:
            path = Path(artifact['path'])
            require(path.is_file() and file_hash(path) == artifact['sha256'], 'artifact_changed', 'Prepared artifact changed or disappeared')
        # Browser task identity is journal metadata, outside the immutable plan.
        if request.get('task_key') and request['task_key'] != state.get('task_key'):
            state['task_key'] = request['task_key']
            atomic_json(directory / 'journal.json', state)
        return state, plan

    def execute(self, request):
        directory = self.directory(request.get('operation_id'))
        with self.lock(directory):
            state, plan = self.bound(directory, request)
            grant = request.get('grant', {})
            case_operation = plan['operation'].startswith('case.')
            if case_operation:
                self.case_journal_boundary()
            target_match = grant.get('targets') == plan['targets'] if case_operation else set(x.get('sku') if isinstance(x, dict) else x for x in grant.get('targets', [])) == set(plan['targets'])
            require(grant.get('execute') is True and grant.get('account') == plan['account'] and grant.get('operation') == plan['operation'] and grant.get('plan_hash') == state['plan_hash'] and target_match, 'grant_mismatch', 'Execution grant must bind exact account, operation, targets and plan hash')
            if plan['body'].get('image_policy') == 'secondary_slots_only' and not state.get('effects_started'):
                try:
                    self.validate_image_baseline(plan, None)
                except OperationError as exc:
                    if exc.code != 'identity_map_incomplete':
                        raise
                    return self.persist(directory, state, 'blocked', reason=exc.code, message=str(exc), effects_started=False, next_action='re-prepare the plan')
            if plan['body'].get('no_changes') and state['status'] == 'prepared':
                return self.persist(directory, state, 'processing', phase='verification', effects_started=False, next_action='reconcile', reason='live_image_verification_required')
            image_preflight_pending = state.get('phase') == 'image_preflight' and state.get('effects_started') is False
            if state['status'] in TERMINAL or state['status'] == 'uncertain' or (state['status'] == 'processing' and not image_preflight_pending) or (state['status'] == 'partial' and not state.get('next_stage') and state.get('next_action') != 'execute'):
                return state
            if case_operation:
                case_tools().validate_execution(plan, grant, case_state_dir=self.case_root, case_policy_path=self.case_policy_path)
            require(state['status'] in {'prepared', 'blocked', 'partial'} or image_preflight_pending, 'not_ready', 'Operation requires inputs before execution')
            adapters = {'flatfilepro.cdp': 'flatfilepro.mjs', 'catalog.cdp': 'catalog.mjs', 'shipment.cdp': 'shipments.mjs', 'cases.cdp': 'cases.mjs'}
            adapter_script = adapters.get(plan['body'].get('adapter'))
            if not adapter_script or not (ROOT / 'tools/amazon-operations' / adapter_script).is_file():
                return self.persist(directory, state, 'blocked', reason='adapter_unavailable', message='Preparation is complete. No verified submission adapter is installed for this operation.')
            require(grant.get('allow_live_canary') is True or grant.get('allow_validated_adapter') is True, 'canary_required', 'This adapter needs an explicitly scoped live canary before production enablement')
            image_preflight = None
            if plan['body'].get('image_policy') == 'secondary_slots_only':
                ffp_source = plan['body'].get('image_catalog_source') == 'flatfilepro_listing_read'
                fresh = self.collect_image_catalog(directory, state, plan, purpose='preflight') if ffp_source else self.collect_export(directory, state, plan, purpose='preflight')
                if fresh is None:
                    return self.persist(directory, state, 'processing', phase='image_preflight', effects_started=False, reason='fresh_catalog_pending', next_action='execute', retry_after_seconds=60)
                try:
                    self.validate_image_baseline(plan, fresh['rows'])
                except OperationError as exc:
                    if exc.code != 'identity_map_incomplete':
                        raise
                    return self.persist(directory, state, 'blocked', reason=exc.code, message=str(exc), effects_started=False, next_action='re-prepare the plan')
                time_field = 'observed_at' if ffp_source else 'report_generated_at'
                require(dt.timedelta(0) <= timestamp(now()) - timestamp(fresh[time_field]) <= dt.timedelta(minutes=5), 'stale_report', 'Image pre-submit evidence must be no more than five minutes old and cannot be future dated')
                image_preflight = {key: fresh[key] for key in ('path', 'sha256', time_field)}
                if ffp_source:
                    image_preflight['source_kind'] = 'flatfilepro_listing_read'
            if plan['body'].get('restore_missing_only'):
                receipt = state.get('preflight_receipt')
                if receipt and timestamp(now()) - timestamp(receipt['report_generated_at']) <= dt.timedelta(minutes=5):
                    require(file_hash(receipt['path']) == receipt['sha256'], 'artifact_changed', 'Fresh preflight report changed')
                    fresh = {**receipt, 'rows': evidence_tools().report_rows(receipt['path'])}
                else:
                    fresh = self.collect_export(directory, state, plan, purpose='preflight')
                if fresh is None:
                    return self.persist(directory, state, 'processing', phase='preflight', effects_started=False, reason='fresh_report_pending', next_action='execute', retry_after_seconds=60)
                self.health_preflight(plan, fresh)
                changes = plan['body'].get('changes', [])
                if all(evidence_tools().field_value(fresh['rows'].get(c['sku'], {}), c['field']) == c['after'] for c in changes):
                    return self.verified_health_noop(directory, state, plan, fresh)
            for image in plan['body'].get('images', []):
                verify_hosted_asset(image['path'], image['url'])
            # Journal intent before spawning; loss of response can never cause blind replay.
            self.persist(directory, state, 'uncertain', execution_started_at=now(), reason='submission_attempt_started', execution_stage=state.pop('next_stage', state.get('execution_stage', 1)), submission_id=None, effects_started=True, phase='executing', next_action=None)
            envelope = {'schema_version': 1, 'plan': plan, 'plan_path': str(directory / 'plan.json'), 'plan_hash': state['plan_hash'], 'receipt_path': str(directory / 'adapter-receipt.json'), 'stage': state['execution_stage'], 'mode': 'execute', 'task_key': state.get('task_key'), 'complete_task': False}
            if image_preflight:
                envelope['image_preflight'] = image_preflight
                envelope['allow_attended_canary'] = grant.get('allow_live_canary') is True
                envelope['allow_validated_image_adapter'] = grant.get('allow_validated_adapter') is True
            atomic_json(directory / 'adapter-input.json', envelope)
            try:
                result = subprocess.run(['node', str(ROOT / 'tools/amazon-operations' / adapter_script), '--request', str(directory / 'adapter-input.json')], capture_output=True, text=True, timeout=1800, check=False, env={**os.environ, **session_environment(os.environ.get('AMAZON_BROWSER_SESSION', 'grimoire'), inherit=True)})
                response = json.loads(result.stdout)
                require(response.get('plan_hash') == state['plan_hash'], 'adapter_identity', 'Adapter result is not bound to this plan')
                status = response.get('status')
                require(status in {'processing', 'blocked', 'failed', 'uncertain'}, 'adapter_status', 'Adapter cannot claim verified completion')
                guarded_submission = case_operation or bool(plan['body'].get('image_policy'))
                if guarded_submission and response.get('attempted') is not False and status in {'blocked', 'failed'}:
                    status = 'uncertain'
                return self.persist(directory, state, status, adapter_result=response, submission_id=response.get('submission_id'), reason=response.get('reason'), effects_started=response.get('attempted', True) if guarded_submission else True)
            except (subprocess.TimeoutExpired, json.JSONDecodeError, OSError, OperationError) as exc:
                return self.persist(directory, state, 'uncertain', reason='adapter_result_uncertain', message=str(exc))

    def validate_image_baseline(self, plan, rows, *, protected_only=False):
        evidence = evidence_tools()
        if not protected_only:
            expected = plan['body'].get('image_identity_rows')
            missing = []
            for sku in plan['targets']:
                identity = expected.get(sku) if isinstance(expected, dict) else None
                if not isinstance(identity, dict) or not image_listing_identity(identity):
                    missing.append(sku)
                    continue
                fields = set(IMAGE_IDENTITY_KEYS)
                if rows is not None:
                    fields.update(image_listing_identity(rows.get(sku, {})))
                fields = sorted(fields - set(identity))
                if fields:
                    missing.append(f'{sku} (missing fields: {", ".join(fields)})')
            legacy = 'Plan predates identity capture; ' if expected is None else ''
            require(not missing and bool(plan['targets']), 'identity_map_incomplete',
                    f'IDENTITY_MAP_INCOMPLETE: {legacy}missing or incomplete identity for SKUs: {", ".join(missing)}; re-prepare the plan')
            if rows is None:
                return
        touched = {(image['sku'], f"other_product_image_locator_{int(image['slot'][2:])}.0.media_location")
                   for image in plan['body']['images'] if image['slot'].startswith('PT')}
        for sku, baseline in plan['body']['image_before_rows'].items():
            for field, value in baseline.items():
                if protected_only and (sku, field) in touched:
                    continue
                observed = evidence.field_value(rows.get(sku, {}), field)
                require(observed is not None and observed == value, 'image_state_conflict', f'Current child identity or protected image changed: {sku}/{field}')
        if not protected_only:
            for sku in plan['targets']:
                require(image_listing_identity(rows.get(sku, {})) == expected[sku], 'image_state_conflict',
                        f'Current listing model, color, size or identity changed: {sku}')

    def record_image_review(self, request):
        """Append an attended review; the journal changes only through reconciliation."""
        directory = self.directory(request.get('operation_id'))
        with self.lock(directory):
            state, plan = self.bound(directory, request)
            require(operation_kind(plan['operation']) == 'listing.images', 'wrong_operation', 'Image operation required')
            require(state.get('submission_id'), 'missing_submission', 'Review requires an existing submission')
            protected_sources = {}
            if any(review.get('purpose') == 'preservation' for review in request['reviews']):
                proof_hash = request.get('review_evidence_digest')
                require(isinstance(proof_hash, str) and bool(re.fullmatch(r'[a-f0-9]{64}', proof_hash)),
                        'review_evidence', 'Protected image review requires its exact collected evidence digest')
                proof_path = directory / f'evidence-{proof_hash}.json'
                proof = json.loads(proof_path.read_text())
                require(digest(proof) == proof_hash and proof.get('account') == plan['account'] and
                        proof.get('plan_hash') == state['plan_hash'] and proof.get('submission_id') == state['submission_id'],
                        'review_evidence', 'Protected review evidence differs from this account, plan or submission')
                expected = self.public_protected_images(plan)
                observations = proof.get('protected_images', [])
                indexed = {(item.get('sku'), item.get('slot')): item for item in observations}
                require(len(indexed) == len(observations), 'review_evidence', 'Duplicate protected image observations')
                for review in request['reviews']:
                    if review.get('purpose') != 'preservation':
                        continue
                    key = (review.get('sku'), review.get('slot'))
                    item = indexed.get(key, {})
                    require(key in expected and item.get('expected_url') == expected[key] and
                            item.get('asin') == plan['body']['sku_asins'][key[0]],
                            'review_evidence', 'Reviewed protected slot or baseline URL differs from the immutable plan')
                    require(timestamp(state.get('execution_started_at', plan['prepared_at'])) <= timestamp(item.get('observed_at')) <= timestamp(now()),
                            'review_evidence', 'Protected image observation must follow submission')
                    require(timestamp(request.get('reviewed_at', now())) >= timestamp(item['observed_at']),
                            'review_evidence', 'Protected image review cannot predate its observed rendition')
                    require(review.get('source_sha256') == item.get('expected_sha256') and
                            review.get('observed_sha256') == item.get('observed_sha256') and
                            review.get('source_id') == item.get('source_id'),
                            'review_evidence', 'Protected review differs from the collected before/after image pair')
                    for field in ('expected_path', 'observed_path'):
                        require(Path(item.get(field, '')).resolve().parent == (directory / 'observed-images').resolve(),
                                'review_evidence', 'Protected review files must be archived in this operation')
                    require(file_hash(item['observed_path']) == item['observed_sha256'], 'review_evidence', 'Collected protected rendition changed')
                    protected_sources[key] = {'sku': key[0], 'slot': key[1], 'path': item['expected_path'],
                        'sha256': item['expected_sha256'], 'purpose': 'preservation', 'baseline_url': expected[key],
                        'baseline_evidence_digest': proof_hash}
            receipt = evidence_tools().build_image_review(plan, request['reviews'], request['reviewer'], request.get('reviewed_at', now()),
                                                        protected_sources=protected_sources)
            require(timestamp(state.get('execution_started_at', plan['prepared_at'])) <= timestamp(receipt['reviewed_at']) <= timestamp(now()),
                    'invalid_time', 'Attended image review must follow submission and cannot be future dated')
            path = directory / f"image-review-{receipt['receipt_sha256']}.json"
            if path.exists():
                require(json.loads(path.read_text()) == receipt, 'review_conflict', 'Immutable review receipt differs')
            else:
                with path.open('x') as stream:
                    stream.write(canonical(receipt).decode() + '\n')
                    stream.flush()
                    os.fsync(stream.fileno())
            return {'status': 'recorded', 'receipt_path': str(path), **receipt}

    def image_release_proof(self, request):
        """Revalidate durable proof with current freshness; never trusts job booleans."""
        directory = self.directory(request.get('operation_id'))
        with self.lock(directory):
            state, plan = self.bound(directory, request)
            pinned = request.get('proof_digest')
            require(pinned is None or bool(re.fullmatch(r'[a-f0-9]{64}', pinned)), 'evidence_mismatch', 'Proof digest must be SHA256')
            path = directory / f'evidence-{pinned}.json' if pinned else Path(state.get('evidence_path') or directory / 'missing-evidence')
            if not path.is_file() or path.resolve().parent != directory.resolve():
                return {'release_eligible': False, 'reason': 'missing_durable_evidence'}
            evidence = json.loads(path.read_text())
            require(path.name == f'evidence-{digest(evidence)}.json', 'evidence_mismatch', 'Durable evidence checksum mismatch')
            reference_time = timestamp(evidence.get('observed_at')) if pinned else None
            require(reference_time is None or reference_time <= timestamp(now()), 'invalid_time', 'Pinned proof cannot be future dated')
            return self.image_completion(directory, state, plan, evidence, reference_time=reference_time)

    def image_completion(self, directory, state, plan, evidence, *, reference_time=None):
        return self._assess_image_completion(directory, state, plan, evidence, reference_time=reference_time)

    def _assess_image_completion(self, directory, state, plan, evidence, *, reference_time=None):
        """Separate publication from contribution processing. Pure except local proof reads."""
        require(operation_kind(plan['operation']) == 'listing.images', 'wrong_operation', 'Image operation required')
        require(evidence.get('account') == plan['account'] and evidence.get('plan_hash') == state['plan_hash'] and
                evidence.get('submission_id') == state.get('submission_id'), 'evidence_mismatch', 'Image evidence identity mismatch')
        body = plan['body']
        clock = reference_time or timestamp(now())
        def fresh(value):
            try:
                observed = timestamp(value)
                return max(timestamp(state.get('execution_started_at', plan['prepared_at'])), clock - dt.timedelta(minutes=15)) <= observed <= clock
            except OperationError:
                return False
        activity = evidence.get('ffp_processing') or {}
        expected_fields = {(sku, evidence_tools().canonical_header(field)): value
                           for sku, fields in body['expected_rows'].items() for field, value in fields.items()}
        attributes = activity.get('attributes', [])
        indexed = {(item.get('sku'), evidence_tools().canonical_header(item.get('field'))): item for item in attributes}
        processing_valid = (activity.get('status') == 'collected' and activity.get('complete') is True and
            activity.get('account') == plan['account'] and activity.get('plan_hash') == state['plan_hash'] and
            activity.get('submission_id') == state.get('submission_id') and fresh(activity.get('observed_at')) and
            len(indexed) == len(attributes) and set(indexed) == set(expected_fields) and
            all(indexed[key].get('submitted_value') == value and indexed[key].get('status') in
                {'reflected', 'pending', 'in_progress', 'failed', 'rejected'} for key, value in expected_fields.items()))
        processing = ('failed' if any(x.get('status') in {'failed', 'rejected'} for x in attributes) else
                      'complete' if all(x.get('status') == 'reflected' for x in attributes) else 'pending') if processing_valid else 'unknown'
        processing_pending = [f'{sku}/{field}' for (sku, field), item in indexed.items()
                              if item.get('status') in {'pending', 'in_progress'}] if processing_valid else []
        preservation = 'unknown'
        receipts = [json.loads(path.read_text()) for path in directory.glob('image-review-*.json')]
        preservation_source = evidence.get('preservation_source') or {}
        if fresh(preservation_source.get('observed_at')):
            try:
                self.validate_image_baseline(plan, evidence.get('protected_rows', {}), protected_only=True)
                preservation = 'verified'
            except OperationError:
                preservation = 'conflict'
        if preservation == 'verified':
            protected_observations = evidence.get('protected_images', [])
            indexed_protected = {(item.get('sku'), item.get('slot')): item for item in protected_observations}
            protected_expected = self.public_protected_images(plan)
            if len(indexed_protected) != len(protected_observations) or set(indexed_protected) != set(protected_expected):
                preservation = 'unknown'
            else:
                for key, expected_url in protected_expected.items():
                    item = indexed_protected[key]
                    try:
                        paths = [Path(item[name]) for name in ('expected_path', 'observed_path')]
                        require(all(path.resolve().parent == (directory / 'observed-images').resolve() for path in paths), 'image_evidence_path', 'Protected image must be archived')
                        before, after = [path.read_bytes() for path in paths]
                        if (item.get('expected_url') != expected_url or item.get('asin') != body['sku_asins'][key[0]] or
                            not fresh(item.get('observed_at')) or hashlib.sha256(before).hexdigest() != item.get('expected_sha256') or
                            hashlib.sha256(after).hexdigest() != item.get('observed_sha256')):
                            preservation = 'unknown'
                        else:
                            source = {'sku': key[0], 'slot': key[1], 'sha256': item['expected_sha256'],
                                      'purpose': 'preservation', 'baseline_url': expected_url}
                            if not evidence_tools().verify_public_rendition(plan, source, before, after, receipts)['content_match']:
                                preservation = 'conflict'
                    except (OSError, ValueError, KeyError):
                        preservation = 'unknown'
        observations = evidence.get('images', [])
        observed = {(item.get('sku'), item.get('slot')): item for item in observations}
        expected_slots = {(image['sku'], image['slot']) for image in body['images']}
        require(len(observed) == len(observations) and set(observed) <= expected_slots, 'image_coverage', 'Unexpected or duplicate image observations')
        matched, pending, failures = [], [], []
        for image in body['images']:
            key = (image['sku'], image['slot'])
            label = '/'.join(key)
            item = observed.get(key, {})
            verified = False
            try:
                observed_path = Path(item.get('observed_path', ''))
                require(observed_path.resolve().parent == (directory / 'observed-images').resolve(), 'image_evidence_path', 'Observed image must be archived in this operation')
                content = observed_path.read_bytes()
                match = evidence_tools().verify_public_rendition(plan, image, Path(image['path']).read_bytes(), content, receipts)
                verified = (bool(match['content_match']) and item.get('observed_sha256') == match['observed_sha256'] and
                    item.get('source_sha256') == image['sha256'] and item.get('asin') == body['sku_asins'][image['sku']] and
                    fresh(item.get('observed_at')) and str(item.get('live_url', '')).startswith('https://') and
                    str(item.get('source_id', '')).startswith('https://'))
            except (OSError, ValueError, KeyError):
                verified = False
            (matched if verified else pending).append(label)
            field = f'other_product_image_locator_{int(image["slot"][2:])}.0.media_location'
            attribute = indexed.get((image['sku'], field), {})
            if processing_valid and attribute.get('status') in {'failed', 'rejected'}:
                failures.append(f'{label}: {attribute["status"]}')
        publication = 'verified' if matched and not pending else ('pending' if observations else 'unobserved')
        eligible = publication == 'verified' and preservation == 'verified' and processing in {'pending', 'complete'}
        return {'publication': publication, 'processing': processing, 'preservation': preservation,
                'release_eligible': eligible, 'proof_digest': digest(evidence), 'observed_at': evidence.get('observed_at'),
                'account': plan['account'], 'operation_id': plan['operation_id'], 'plan_hash': state['plan_hash'],
                'submission_id': state.get('submission_id'), 'matched': matched, 'pending': pending, 'failures': failures,
                'processing_pending': sorted(processing_pending), 'freshness_seconds': 900}

    def refresh_image_observations(self, request, *, persist=False):
        """Rebuild display history from local proofs only; never run a collector."""
        directory = self.directory(request.get('operation_id'))
        with self.lock(directory):
            state, plan = self.bound(directory, request)
            self._update_image_observations(directory, state, plan)
            if persist:
                atomic_json(directory / 'journal.json', state)
            return state

    def _update_image_observations(self, directory, state, plan, evidence=None):
        """Compact, scope-bound reporting facts; completion never consumes these."""
        binding = {key: state.get(key) for key in ('operation_id', 'plan_hash', 'submission_id')}
        binding['account'] = plan['account']
        history = state.get('image_observations')
        rebuild = not isinstance(history, dict) or history.get('binding') != binding
        if rebuild:
            history = {'schema_version': 1, 'binding': binding,
                       **{name: {'slots': {}} if name == 'publication' else {}
                          for name in ('publication', 'processing', 'preservation')}}
        samples = []
        if rebuild:
            for path in directory.glob('evidence-*.json'):
                try:
                    value = json.loads(path.read_text())
                    if (isinstance(value, dict) and path.name == f'evidence-{digest(value)}.json' and
                            value.get('account') == plan['account'] and value.get('plan_hash') == state['plan_hash'] and
                            value.get('submission_id') == state.get('submission_id') and
                            timestamp(state.get('execution_started_at', plan['prepared_at'])) <= timestamp(value.get('observed_at')) <= timestamp(now())):
                        samples.append(value)
                except (OSError, ValueError):
                    continue
            samples.sort(key=lambda value: timestamp(value['observed_at']))
        if evidence is not None and (not samples or digest(samples[-1]) != digest(evidence)):
            samples.append(evidence)
        samples.sort(key=lambda value: timestamp(value['observed_at']), reverse=True)
        expected = {(row['sku'], row['slot']): row for row in plan['body']['images']}

        def remember(name, value, dates, proof, *, slots=None):
            if not dates:
                return
            observed = min(dates, key=timestamp)
            fact = {'value': value, 'observed_at': observed, 'evidence_digest': proof}
            if slots is not None:
                fact['matched'] = slots
            old = history[name].get('last_successful')
            if not old or timestamp(observed) >= timestamp(old['observed_at']):
                history[name]['last_successful'] = fact
            if value == 'verified':
                old = history[name].get('last_verified')
                if not old or timestamp(observed) >= timestamp(old['observed_at']):
                    history[name]['last_verified'] = fact

        publication_coverage = {}
        for sample in samples:
            if rebuild and all(history[name].get('last_successful') for name in ('publication', 'processing', 'preservation')):
                break
            try:
                activity = sample.get('ffp_processing') or {}
                if not isinstance(activity, dict):
                    continue
                if rebuild and not any((
                        not history['publication'].get('last_successful') and sample.get('images') and not sample.get('pdp_collection_error'),
                        not history['processing'].get('last_successful') and activity.get('status') == 'collected',
                        not history['preservation'].get('last_successful') and sample.get('preservation_source'))):
                    continue
                dates = [sample['observed_at']]
                dates += [r['observed_at'] for r in sample.get('images', []) + sample.get('protected_images', []) if r.get('observed_at')]
                dates += [sample[k]['observed_at'] for k in ('ffp_processing', 'preservation_source')
                          if isinstance(sample.get(k), dict) and sample[k].get('observed_at')]
                moment = max(map(timestamp, dates))
                if moment > timestamp(now()):
                    continue
                # Reuse the execution validator at the original observation time.
                # This validates historical facts without refreshing their age.
                completion = self._assess_image_completion(directory, state, plan, sample, reference_time=moment)
                proof = digest(sample)
                activity = sample.get('ffp_processing') or {}
                if completion['processing'] != 'unknown':
                    remember('processing', completion['processing'], [activity['observed_at']], proof)
                protected = sample.get('preservation_source') or {}
                if completion['preservation'] != 'unknown':
                    dates = [protected['observed_at']] + [row['observed_at'] for row in sample.get('protected_images', [])]
                    remember('preservation', completion['preservation'], dates, proof)
                valid = []
                for row in sample.get('images', []):
                    key = (row.get('sku'), row.get('slot'))
                    source = expected.get(key)
                    if source is None or sample.get('pdp_collection_error'):
                        continue
                    observed = timestamp(row.get('observed_at'))
                    path = Path(row.get('observed_path', ''))
                    if (not moment - dt.timedelta(seconds=900) <= observed <= moment or
                            observed < timestamp(state.get('execution_started_at', plan['prepared_at'])) or
                            path.resolve().parent != (directory / 'observed-images').resolve() or
                            file_hash(path) != row.get('observed_sha256') or
                            row.get('source_sha256') != source['sha256'] or
                            file_hash(source['path']) != source['sha256'] or
                            row.get('asin') != plan['body']['sku_asins'][key[0]] or
                            not str(row.get('source_id', '')).startswith('https://') or
                            not str(row.get('live_url', '')).startswith('https://')):
                        continue
                    label = '/'.join(key)
                    fact = {'value': 'verified' if label in completion['matched'] else 'unverified',
                            'observed_at': row['observed_at'], 'evidence_digest': proof}
                    old = history['publication']['slots'].get(label)
                    if not old or observed >= timestamp(old['observed_at']):
                        history['publication']['slots'][label] = fact
                    valid.append(row)
                if len(valid) == len(expected):
                    remember('publication', completion['publication'], [row['observed_at'] for row in valid],
                             proof, slots=completion['matched'])
                publication_coverage[proof] = len(valid) == len(sample.get('images', []))
            except (OSError, ValueError, KeyError, TypeError, AttributeError):
                continue  # Bad history cannot establish a reporting fact.
        if evidence is None and samples:
            evidence = samples[0]
        if evidence is not None:
            sources = {'publication': evidence.get('pdp_collection_error') or evidence.get('pdp_collection'),
                       'processing': evidence.get('ffp_processing'),
                       'preservation': evidence.get('preservation_source') or state.get('last_catalog_collection')}
            attempts = []
            for name, response in sources.items():
                response = response or {}
                if not isinstance(response, dict):
                    response = {'status': 'invalid', 'reason': 'Malformed collection result.'}
                if state.get('evidence_skipped') and name != 'processing':
                    attempts.append(history[name].get('last_attempt') or {'outcome': 'collected', 'retryable': False, 'reason': ''})
                    continue
                detail = ' '.join(str(response.get(k) or '') for k in ('code', 'reason', 'message'))
                transient = bool(re.search(r'BROWSER_SESSION_BUSY|ECONNRESET|ETIMEDOUT|EAI_AGAIN|HTTP (?:429|502|503|504)|timed?\s*out|timeout|collector_result_unavailable', detail, re.I))
                outcome = response.get('status', 'unavailable')
                if name == 'publication' and outcome == 'collected' and not publication_coverage.get(digest(evidence), True):
                    outcome = 'invalid'
                    detail = 'Image evidence lacks verified provenance or complete coverage.'
                if outcome not in {'collected', 'partial', 'blocked', 'invalid'}:
                    outcome = 'unavailable'
                attempt = {'attempted_at': evidence['observed_at'], 'outcome': outcome,
                           'retryable': transient, 'reason': detail.strip()[:500]}
                history[name]['last_attempt'] = attempt
                attempts.append(attempt)
            state['image_collection_attempt'] = {
                'attempted_at': evidence['observed_at'],
                'outcome': 'deferred' if any(a['retryable'] for a in attempts) else
                           'collected' if all(a['outcome'] == 'collected' for a in attempts) else 'incomplete',
                'reasons': list(dict.fromkeys(a['reason'] for a in attempts if a['reason']))}
        history['reconstructed'] = True
        state['image_observations'] = history
        return history

    def public_protected_images(self, plan):
        """Public MAIN and populated untouched secondary slots; swatches retain catalog proof."""
        touched = {(image['sku'], image['slot']) for image in plan['body']['images']}
        expected = {}
        for sku, row in plan['body'].get('image_before_rows', {}).items():
            for field, url in row.items():
                canonical_field = evidence_tools().canonical_header(field)
                match = re.fullmatch(r'other_product_image_locator_([1-9])\.0\.media_location', canonical_field)
                slot = 'MAIN' if canonical_field == 'main_product_image_locator.0.media_location' else f'PT0{match[1]}' if match else None
                if slot and url and (sku, slot) not in touched:
                    expected[(sku, slot)] = url
        return expected

    def verified_health_noop(self, directory, state, plan, fresh):
        evidence = {'schema_version': 1, 'account': plan['account'], 'plan_hash': state['plan_hash'], 'source_id': fresh.get('source_id', fresh['path']), 'observed_at': now(), 'processing_status': 'live_observed', 'no_changes': True, 'report_artifact': {key: fresh[key] for key in ('path', 'sha256', 'report_generated_at') if key in fresh}, 'rows': {sku: {field: evidence_tools().field_value(fresh['rows'].get(sku, {}), field) for field in fields} for sku, fields in plan['body'].get('expected_rows', {}).items()}}
        path = directory / f'evidence-{digest(evidence)}.json'
        atomic_json(path, evidence)
        return self.persist(directory, state, 'verified', no_changes=True, effects_started=False, reason='already_restored', evidence_path=str(path))

    def health_preflight(self, plan, fresh):
        changes = plan['body'].get('changes', [])
        evidence = evidence_tools()
        values = [evidence.field_value(fresh['rows'].get(c['sku'], {}), c['field']) for c in changes]
        if changes and all(value == change['after'] for value, change in zip(values, changes)):
            return
        require(changes and all(value is not None and not value.strip() for value in values), 'health_conflict', 'Fresh live fields are no longer all missing; automatic restoration stopped')
        changed_cells = {(c['sku'], evidence.canonical_header(c['field'])) for c in changes}
        for sku, fields in plan['body'].get('expected_rows', {}).items():
            for field, expected in fields.items():
                if (sku, evidence.canonical_header(field)) in changed_cells:
                    continue
                current = evidence.field_value(fresh['rows'].get(sku, {}), field)
                require(current is not None and current == as_text(expected), 'health_conflict', f'Carried field {sku}/{field} is missing or changed; automatic restoration stopped')

    def run_collector(self, directory, script, request, timeout=180):
        require(script in {'catalog-export.mjs', 'image-evidence.mjs', 'cases.mjs', 'flatfilepro.mjs', 'flatfilepro-export.mjs', 'flatfilepro-activity.mjs', 'flatfilepro-listings.mjs', 'flatfilepro-discovery.mjs'}, 'invalid_collector', 'Unknown fixed collector')
        if _CLOSE_RECONCILE_TABS.get():
            request = {**request, 'close_tab_after': True}
        input_path = directory / (script.replace('.mjs', '') + '-input.json')
        atomic_json(input_path, request)
        try:
            process = subprocess.Popen(['node', str(ROOT / 'tools/amazon-operations' / script), '--request', str(input_path)], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, env={**os.environ, **session_environment(os.environ.get('AMAZON_BROWSER_SESSION', 'grimoire'), inherit=True)})
            try:
                stdout, _ = process.communicate(timeout=timeout)
            except subprocess.TimeoutExpired:
                process.terminate()
                try:
                    process.communicate(timeout=15)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.communicate()
                return {'status': 'blocked', 'reason': 'collector_result_unavailable'}
            data = json.loads(stdout)
            if data.get('status') == 'collected':
                require(data.get('account') == request['account'] and data.get('plan_hash') == request['plan_hash'], 'collector_identity', 'Collector account or plan mismatch')
            return data
        except (OSError, ValueError, subprocess.TimeoutExpired):
            return {'status': 'blocked', 'reason': 'collector_result_unavailable'}

    def observe(self, request):
        """Collect case capability and correspondence without preparing a send."""
        require(request.get('schema_version') == 1 and request.get('operation') in {'case.create', 'case.reply'}, 'invalid_request', 'Case observation requires a case operation')
        account(request.get('account'))
        case_tools().validate_targets(request['operation'], request.get('targets'))
        directory = self.directory(request.get('operation_id'))
        return self.run_collector(directory, 'cases.mjs', {**request, 'mode': 'observe', 'plan_hash': digest({k: v for k, v in request.items() if k not in {'task_key', 'complete_task'}})})

    def collect_export(self, directory, state, plan, purpose='verification', *, complete_task=False):
        minimum = plan['prepared_at'] if purpose == 'preflight' else state.get('execution_started_at', plan['prepared_at'])
        request = {'schema_version': 1, 'operation_id': plan['operation_id'], 'account': plan['account'], 'plan_hash': state['plan_hash'], 'minimum_after': minimum, 'output_dir': str(directory / f"backend-{purpose}-{state.get('execution_stage', 1)}"), 'task_key': state.get('task_key'), 'complete_task': complete_task is True}
        response = self.run_collector(directory, 'catalog-export.mjs', request)
        state['last_collection'] = {k: v for k, v in response.items() if k not in {'account', 'plan_hash'}}
        if response.get('status') != 'collected':
            return None
        require(response.get('complete_report') is True and timestamp(response.get('report_generated_at')) >= timestamp(minimum), 'stale_report', 'Verification requires a complete freshly generated report')
        path = Path(response['path'])
        require(path.is_file() and file_hash(path) == response.get('sha256'), 'artifact_changed', 'Downloaded report checksum mismatch')
        response['rows'] = evidence_tools().report_rows(path)
        return response

    def collect_image_catalog(self, directory, state, plan, purpose='verification', *, complete_task=False):
        minimum = plan['prepared_at'] if purpose == 'preflight' else state.get('execution_started_at', plan['prepared_at'])
        response = self.run_collector(directory, 'flatfilepro-listings.mjs', {
            'schema_version': 1, 'operation_id': plan['operation_id'], 'account': plan['account'], 'plan_hash': state['plan_hash'],
            'targets': [{'sku': sku, 'asin': plan['body']['sku_asins'][sku]} for sku in plan['body']['image_before_rows']],
            'minimum_after': minimum, 'output_dir': str(directory / f'ffp-{purpose}'),
            'task_key': state.get('task_key'), 'complete_task': complete_task is True})
        state['last_catalog_collection'] = response
        if response.get('status') != 'collected':
            return None
        require(response.get('complete') is True and response.get('source_kind') == 'flatfilepro_listing_read' and
                timestamp(minimum) <= timestamp(response.get('observed_at')) <= timestamp(now()), 'stale_report', 'Complete current FlatFilePro read required')
        require(file_hash(response['path']) == response.get('sha256'), 'artifact_changed', 'FlatFilePro listing evidence changed')
        saved = json.loads(Path(response['path']).read_text())
        require(saved.get('rows') == response.get('rows') and saved.get('account') == plan['account'] and saved.get('plan_hash') == state['plan_hash'], 'artifact_changed', 'Listing evidence differs from receipt')
        require(set(response.get('rows', {})) == set(plan['body']['image_before_rows']), 'scope_violation', 'Listing read SKU coverage differs')
        return response

    def collect_images(self, directory, state, plan, *, complete_task=False, full_verify_seconds=IMAGE_FULL_VERIFY_SECONDS):
        body = plan['body']
        require(isinstance(full_verify_seconds, (int, float)) and not isinstance(full_verify_seconds, bool) and
                math.isfinite(full_verify_seconds) and full_verify_seconds >= 0,
                'invalid_full_verify_seconds', 'full_verify_seconds must be a finite nonnegative number')
        state.pop('evidence_skipped', None)
        activity = None
        previous_activity = state.get('last_processing_collection')
        if body.get('image_policy') and state.get('submission_id'):
            activity = self.run_collector(directory, 'flatfilepro-activity.mjs', {
                'schema_version': 1, 'operation_id': plan['operation_id'], 'account': plan['account'],
                'plan_hash': state['plan_hash'], 'submission_id': state['submission_id'],
                'expected_rows': body['expected_rows'], 'task_key': state.get('task_key'), 'complete_task': False}, timeout=300)
            state['last_processing_collection'] = activity
            # Read timestamps change on every poll; all other Activity values,
            # including counts and failures, must agree before reusing evidence.
            def snapshot(value):
                if not isinstance(value, dict):
                    return None
                return {**{key: item for key, item in value.items() if key not in {'observed_at', 'attributes'}},
                        'attributes': sorted(value.get('attributes', []), key=canonical)}
            attributes = activity.get('attributes', [])
            summary = activity.get('summary') or {}
            pending = (not body.get('no_changes') and activity.get('status') == 'collected' and activity.get('complete') is True and
                       activity.get('account') == plan['account'] and activity.get('plan_hash') == state['plan_hash'] and
                       activity.get('submission_id') == state['submission_id'] and
                       any(item.get('status') in {'pending', 'in_progress'} for item in attributes) and
                       all(item.get('status') in {'pending', 'in_progress', 'reflected'} for item in attributes) and
                       not any(summary.get(key) or activity.get(key) for key in ('failed', 'rejected', 'errors', 'failures')) and
                       activity.get('processing_status', 'processing') == 'processing')
            if (pending and complete_task is not True and state.get('status') in {'processing', 'partial'} and
                    state.get('image_collection_attempt', {}).get('outcome') not in {'deferred', 'incomplete'} and
                    isinstance(state.get('image_completion'), dict) and snapshot(activity) == snapshot(previous_activity)):
                try:
                    age = (timestamp(now()) - timestamp(state.get('last_image_full_read_at'))).total_seconds()
                    path = Path(state.get('evidence_path') or '')
                    if 0 <= age < full_verify_seconds and path.resolve().parent == directory.resolve():
                        saved = json.loads(path.read_text())
                        if (isinstance(saved, dict) and path.name == f'evidence-{digest(saved)}.json' and
                                saved.get('account') == plan['account'] and saved.get('plan_hash') == state['plan_hash'] and
                                saved.get('submission_id') == state['submission_id'] and
                                saved.get('processing_status') == 'live_observed' and
                                snapshot(saved.get('ffp_processing')) == snapshot(activity)):
                            state['evidence_skipped'] = 'activity-unchanged'
                            # Preserve every PDP and preservation observation time.
                            return {**saved, 'ffp_processing': activity, 'observed_at': now()}
                except (OSError, ValueError):
                    pass
        mapping = body.get('sku_asins', {})
        if not all(image['sku'] in mapping for image in body['images']):
            return None
        targets = {}
        for image in body['images']:
            targets.setdefault(image['sku'], {'asin': mapping[image['sku']], 'slots': []})['slots'].append(image['slot'])
        protected_expected = self.public_protected_images(plan) if body.get('image_policy') else {}
        for sku, slot in protected_expected:
            targets[sku]['slots'].append(slot)
        previous = state.get('last_collection')
        response = self.run_collector(directory, 'image-evidence.mjs', {'schema_version': 1, 'operation_id': plan['operation_id'], 'account': plan['account'], 'plan_hash': state['plan_hash'], 'targets': targets,
            'previous': previous, 'task_key': state.get('task_key'), 'complete_task': False}, timeout=300)
        collection_error = None
        if (response.get('status') == 'blocked' and previous and previous.get('status') in {'partial', 'collected'} and
                previous.get('account') == plan['account'] and previous.get('plan_hash') == state['plan_hash'] and
                previous.get('operation_id') == plan['operation_id']):
            # Keep prior PDP observations with their original timestamps. A new
            # failed read does not refresh them or conceal the current failure.
            state['last_collection_error'] = response
            collection_error = response
            response = previous
        else:
            state['last_collection'] = response
            state.pop('last_collection_error', None)
        approved = {(image['sku'], image['slot']): image for image in body['images']}
        observed = []
        receipts = [json.loads(path.read_text()) for path in directory.glob('image-review-*.json')]
        archives = directory / 'observed-images'
        archives.mkdir(exist_ok=True)
        downloads = {}
        protected_images = []
        for image in response.get('images', []):
            source = approved.get((image['sku'], image['slot']))
            protected_url = protected_expected.get((image['sku'], image['slot']))
            require(source is not None or protected_url, 'scope_violation', 'Image collector returned an unrequested slot')
            try:
                if image['live_url'] not in downloads:
                    downloads[image['live_url']] = evidence_tools().fetch_public_image(image['live_url'])
                content = downloads[image['live_url']]
                if protected_url:
                    if protected_url not in downloads:
                        downloads[protected_url] = evidence_tools().fetch_public_image(protected_url)
                    before = downloads[protected_url]
                    pair = {**image, 'expected_url': protected_url}
                    for label, data in [('expected', before), ('observed', content)]:
                        checksum = hashlib.sha256(data).hexdigest()
                        path = archives / checksum
                        if not path.exists():
                            with path.open('xb') as stream:
                                stream.write(data)
                        pair.update({label + '_path': str(path), label + '_sha256': checksum})
                    protected_images.append(pair)
                    continue
                match = evidence_tools().verify_public_rendition(plan, source, Path(source['path']).read_bytes(), content, receipts)
                path = archives / match['observed_sha256']
                if path.exists():
                    require(file_hash(path) == match['observed_sha256'], 'image_evidence_changed', 'Archived rendition changed')
                else:
                    with path.open('xb') as stream:
                        stream.write(content)
                observed.append({**image, 'source_sha256': source['sha256'], 'observed_path': str(path),
                    'observed_at': image.get('observed_at', response.get('observed_at')),
                    'visually_verified': bool(match['content_match']), **match})
            except (OSError, ValueError) as exc:
                if source:
                    observed.append({**image, 'source_sha256': source['sha256'], 'visually_verified': False, 'content_match': None, 'error': str(exc)})
                else:
                    protected_images.append({**image, 'error': str(exc)})
        protected = {}
        if body.get('image_policy'):
            try:
                ffp_source = (body.get('image_catalog_source') == 'flatfilepro_listing_read' or
                              (body.get('image_policy') == 'secondary_slots_only' and body.get('adapter') == 'flatfilepro.cdp'))
                fresh = self.collect_image_catalog(directory, state, plan, complete_task=False) if ffp_source else self.collect_export(directory, state, plan, complete_task=False)
                if fresh is not None:
                    protected = {'protected_rows': fresh['rows'], 'preservation_source':
                        {key: value for key, value in fresh.items() if key != 'rows'}}
                    # A report timestamp remains a report timestamp, not a new read time.
                    protected['preservation_source'].setdefault('observed_at', fresh.get('report_generated_at'))
            except (OSError, ValueError) as exc:
                protected = {'preservation_source': {'status': 'blocked', 'message': str(exc)}}
        state['last_image_full_read_at'] = now()
        return {'account': plan['account'], 'plan_hash': state['plan_hash'], 'source_id': 'amazon-live-imageblock', 'observed_at': now(), 'submission_id': state.get('submission_id'), 'processing_status': 'live_observed', 'images': observed, **protected,
                'pdp_collection': response, 'protected_images': protected_images,
                **({'pdp_collection_error': collection_error} if collection_error else {}),
                **({'ffp_processing': activity} if activity else {})}

    def collect(self, directory, state, plan, *, complete_task=False):
        """Read-only live PDP collector using the established listing-capture runner.

        Public title/bullet/description fields can be verified here. Hidden fields
        and seller contributions require fresh export evidence; never substitute
        staged import values for live values.
        """
        if state['status'] in TERMINAL:
            return None
        body = plan['body']
        if plan['operation'].startswith('case.'):
            return case_tools().collect(self, directory, state, plan)
        driver_evidence = state.get('adapter_result', {}).get('evidence')
        if isinstance(driver_evidence, dict):
            return driver_evidence
        kind = operation_kind(plan['operation'])
        if kind == 'listing.images' and state.get('effects_started') and not state.get('submission_id'):
            recovery = self.run_collector(directory, 'flatfilepro.mjs', {'schema_version': 1, 'plan': plan, 'plan_path': str(directory / 'plan.json'), 'plan_hash': state['plan_hash'], 'receipt_path': str(directory / 'adapter-receipt.json'), 'account': plan['account'], 'mode': 'reconcile', 'task_key': state.get('task_key'), 'complete_task': False})
            state['last_import_recovery'] = recovery
            if recovery.get('status') == 'collected' and recovery.get('submission_id'):
                state['submission_id'] = recovery['submission_id']
                if recovery.get('processing_status') == 'failed':
                    return recovery
        if kind == 'listing.images' and (state.get('submission_id') or body.get('no_changes')):
            return self.collect_images(directory, state, plan, complete_task=complete_task)
        if not state.get('submission_id'):
            return None
        if kind == 'catalog.change':
            exported = self.collect_export(directory, state, plan, complete_task=complete_task)
            if exported is None:
                return None
            return {'account': plan['account'], 'plan_hash': state['plan_hash'], 'source_id': exported['source_id'], 'observed_at': now(), 'submission_id': state['submission_id'], 'processing_status': 'live_observed', **evidence_tools().catalog_evidence(plan, state, exported['rows'])}
        if kind not in {'seo.update', 'flatfilepro.update'}:
            return None
        expected = body.get('expected_rows', {})
        hidden_fields = any(not re.fullmatch(r'item_name.0.value|itemName|bullet_point.[0-9]+.value|product_description.0.value', field) for row in expected.values() for field in row)
        if hidden_fields or not body.get('sku_asins'):
            exported = self.collect_export(directory, state, plan, complete_task=complete_task)
            if exported is None:
                return None
            rows = {sku: {field: value for field in fields if (value := evidence_tools().field_value(exported['rows'].get(sku, {}), field)) is not None} for sku, fields in expected.items()}
            return {'account': plan['account'], 'plan_hash': state['plan_hash'], 'source_id': exported['source_id'], 'observed_at': now(), 'submission_id': state['submission_id'], 'processing_status': 'live_observed', 'rows': rows}
        mapping = body.get('sku_asins', {})
        expected = body.get('expected_rows', {})
        if not expected or not set(expected) <= set(mapping):
            return None
        require(all(re.fullmatch(r'[A-Z0-9]{10}', mapping[sku]) for sku in expected), 'invalid_asin', 'Live collector requires valid ASINs')
        market = {'US': ('com', 'en_US'), 'DE': ('de', 'de_DE'), 'AU': ('com.au', 'en_AU'), 'UK': ('co.uk', 'en_GB'), 'IT': ('it', 'it_IT'), 'FR': ('fr', 'fr_FR'), 'ES': ('es', 'es_ES'), 'CA': ('ca', 'en_CA')}.get(plan['account']['marketplace'])
        if not market:
            return None
        output = directory / ('live-listings-' + str(len(state.get('events', []))) + '.json')
        try:
            result = subprocess.run(['node', str(ROOT / 'tools/listing-capture/capture-cdp.mjs'), ','.join(sorted(set(mapping[sku] for sku in expected))), str(output), *market], capture_output=True, text=True, timeout=300, check=False, env={**os.environ, **session_environment(os.environ.get('AMAZON_BROWSER_SESSION', 'grimoire'), inherit=True)})
            if result.returncode or not output.is_file():
                return None
            captured = json.loads(output.read_text())
            listings = {x['asin']: x for x in captured.get('listings', []) if x.get('status') == 'ok' and x.get('resolvedAsin') == x.get('asin')}
            rows = {}
            for sku, fields in expected.items():
                listing = listings.get(mapping[sku])
                if not listing:
                    continue
                rows[sku] = {}
                for field in fields:
                    if field in {'itemName', 'item_name.0.value'} and listing.get('title'):
                        rows[sku][field] = listing['title']
                    elif field == 'product_description.0.value' and listing.get('description'):
                        rows[sku][field] = listing['description']
                    elif re.fullmatch(r'bullet_point.[0-9]+.value', field):
                        index = int(field.split('.')[1])
                        if index < len(listing.get('bullets', [])):
                            rows[sku][field] = listing['bullets'][index]
            return {'account': plan['account'], 'plan_hash': state['plan_hash'], 'source_id': str(output), 'observed_at': now(), 'submission_id': state['submission_id'], 'processing_status': 'live_observed', 'rows': rows, 'collection': 'Live retail content; feed processing status is not inferred'}
        except (OSError, subprocess.TimeoutExpired, ValueError, KeyError):
            return None

    def _complete_reconciled_task(self, request, state, plan, status):
        if request.get('complete_task') is not True or status not in {'verified', 'failed', 'blocked', 'stalled'}:
            return
        try:
            # Match taskIdFor('amazon-operations', task_key || operation_id).
            key = state.get('task_key') or plan['operation_id']
            task_id = 'amazon-operations:' + hashlib.sha256(str(key).encode()).hexdigest()[:20]
            env = {**os.environ, **session_environment(os.environ.get('AMAZON_BROWSER_SESSION', 'grimoire'), inherit=True)}
            subprocess.run(['node', str(ROOT / 'tools/browserctl/browserctl.mjs'), 'task', 'complete',
                            '--port', env['CDP_PORT'], '--task-id', task_id],
                           capture_output=True, text=True, timeout=30, check=True, env=env)
        except Exception as exc:
            print(f'Browser task completion failed: {exc}', file=sys.stderr)

    def reconcile(self, request):
        """close_tabs controls collector envelopes; finalize='stalled' preserves evidence."""
        token = _CLOSE_RECONCILE_TABS.set(request.get('close_tabs') is True)
        try:
            return self._reconcile(request)
        finally:
            _CLOSE_RECONCILE_TABS.reset(token)

    def _reconcile(self, request):
        directory = self.directory(request.get('operation_id'))
        with self.lock(directory):
            state, plan = self.bound(directory, request)
            require('finalize' not in request or request['finalize'] == 'stalled',
                    'invalid_finalize', 'finalize must be stalled when supplied')
            if request.get('finalize') == 'stalled':
                if state['status'] == 'stalled':
                    return state
                result = self.persist(directory, state, 'stalled', stalled_from_status=state['status'],
                    journal_note='Finalized as stalled by request; last evidence retained.')
                self._complete_reconciled_task(request, state, plan, 'stalled')
                return result
            if state['status'] == 'stalled':
                state['status'] = state.pop('stalled_from_status', 'processing')
                state.pop('journal_note', None)
            if plan['operation'].startswith('case.'):
                self.case_journal_boundary()
                return case_tools().reconcile(self, directory, state, plan, request)
            if state.get('phase') == 'image_preflight' and state.get('effects_started') is False:
                return self.persist(directory, state, 'processing', next_action='execute', reason='fresh_report_pending', retry_after_seconds=60)
            if state.get('phase') == 'preflight' and state.get('effects_started') is False:
                require(plan['body'].get('restore_missing_only'), 'invalid_preflight', 'Unexpected preflight state')
                fresh = self.collect_export(directory, state, plan, purpose='preflight')
                if fresh is None:
                    return self.persist(directory, state, 'processing', phase='preflight', effects_started=False, next_action='execute', reason='fresh_report_pending', retry_after_seconds=60)
                self.health_preflight(plan, fresh)
                changes = plan['body']['changes']
                if all(evidence_tools().field_value(fresh['rows'].get(c['sku'], {}), c['field']) == c['after'] for c in changes):
                    self._complete_reconciled_task(request, state, plan, 'verified')
                    return self.verified_health_noop(directory, state, plan, fresh)
                receipt = {key: fresh[key] for key in ('path', 'sha256', 'report_generated_at', 'observed_at')}
                return self.persist(directory, state, 'partial', phase='preflight', effects_started=False, next_action='execute', reason='fresh_missing_fields_verified', preflight_receipt=receipt)
            evidence = request.get('evidence')
            state.pop('evidence_skipped', None)
            if not evidence:
                if (operation_kind(plan['operation']) == 'listing.images' and state.get('submission_id') and
                        state['status'] not in TERMINAL and not isinstance(state.get('adapter_result', {}).get('evidence'), dict)):
                    evidence = self.collect_images(directory, state, plan, complete_task=request.get('complete_task') is True,
                        full_verify_seconds=request.get('full_verify_seconds', IMAGE_FULL_VERIFY_SECONDS))
                else:
                    evidence = self.collect(directory, state, plan, complete_task=False)
                if evidence is None:
                    self._complete_reconciled_task(request, state, plan, state['status'])
                    return self.persist(directory, state, state['status'], evidence_required=['Fresh operation-scoped Amazon processing result and final live state'], reason='verification_evidence_required')
            scoped(evidence, plan['account'], 'reconciliation evidence')
            require(evidence.get('plan_hash') == state['plan_hash'] and evidence.get('source_id'), 'evidence_mismatch', 'Evidence requires exact plan_hash and source_id')
            require(timestamp(evidence.get('observed_at')) >= timestamp(state.get('execution_started_at', plan['prepared_at'])), 'stale_evidence', 'Evidence predates this operation')
            require(timestamp(evidence['observed_at']) <= timestamp(now()) + dt.timedelta(minutes=5), 'invalid_time', 'Evidence cannot be future dated')
            if state.get('submission_id'):
                require(evidence.get('submission_id') == state['submission_id'], 'submission_mismatch', 'Evidence submission ID mismatch')
            op, body = operation_kind(plan['operation']), plan['body']
            require(body.get('status') not in {'awaiting_input'}, 'not_ready', 'Cannot reconcile an incomplete plan')
            if state['status'] == 'verified':
                return state
            require(evidence.get('submission_id') or body.get('no_changes'), 'missing_submission', 'Submission evidence requires its external ID')
            processing = evidence.get('processing_status')
            require(processing in {'processing', 'complete', 'failed', 'live_observed'}, 'invalid_processing', 'processing_status must be processing, complete, failed, or live_observed')
            evidence_path = directory / f'evidence-{digest(evidence)}.json'
            atomic_json(evidence_path, evidence)
            if processing == 'processing':
                return self.persist(directory, state, 'processing', submission_id=evidence['submission_id'], evidence_path=str(evidence_path))
            if processing == 'failed':
                self._complete_reconciled_task(request, state, plan, 'failed')
                return self.persist(directory, state, 'failed', submission_id=evidence['submission_id'], evidence_path=str(evidence_path), failures=evidence.get('errors', ['Submission failed']))
            if op == 'listing.images' and body.get('image_policy') and not body.get('no_changes'):
                while True:
                    completion = self.image_completion(directory, state, plan, evidence)
                    status = ('partial' if completion['processing_pending'] and
                              (completion['processing'] == 'failed' or completion['preservation'] == 'conflict') else
                              'failed' if completion['processing'] == 'failed' or completion['preservation'] == 'conflict' else
                              'verified' if completion['release_eligible'] and completion['processing'] == 'complete' else
                              'partial' if completion['matched'] else 'processing')
                    if not state.get('evidence_skipped') or status not in {'verified', 'failed', 'blocked'}:
                        break
                    # A terminal decision must use collectors from this pass.
                    evidence = self.collect_images(directory, state, plan, complete_task=True)
                    require(evidence is not None, 'missing_image_evidence', 'Terminal image verification requires current reads')
                    evidence_path = directory / f'evidence-{digest(evidence)}.json'
                    atomic_json(evidence_path, evidence)
                self._update_image_observations(directory, state, plan, evidence)
                if state.get('evidence_skipped'):
                    # Unchanged Activity carries the persisted publication result
                    # forward. Freshness above still gates terminal decisions.
                    return self.persist(directory, state, state['status'], evidence_path=str(evidence_path))
                self._complete_reconciled_task(request, state, plan, status)
                return self.persist(directory, state, status, submission_id=evidence.get('submission_id'),
                    matched=completion['matched'], pending=completion['pending'], failures=completion['failures'],
                    evidence_path=str(evidence_path), image_completion=completion)
            matched, pending = [], []
            if op in {'seo.update', 'flatfilepro.update'}:
                observed = evidence.get('rows', {})
                for sku, row in body.get('expected_rows', {}).items():
                    for field, value in row.items():
                        (matched if sku in observed and field in observed[sku] and as_text(observed[sku][field]) == value else pending).append(f'{sku}/{field}')
            elif op == 'listing.images':
                if body.get('image_policy'):
                    self.validate_image_baseline(plan, evidence.get('protected_rows', {}), protected_only=True)
                processing_attributes = {}
                if body.get('image_policy') and not body.get('no_changes'):
                    activity = evidence.get('ffp_processing') or {}
                    require(activity.get('status') == 'collected' and activity.get('complete') is True and
                            activity.get('submission_id') == evidence.get('submission_id') and
                            activity.get('account') == plan['account'] and activity.get('plan_hash') == state['plan_hash'],
                            'missing_processing_evidence', 'Secondary image verification requires complete exact-run FlatFilePro processing evidence')
                    processing_attributes = {(item['sku'], evidence_tools().canonical_header(item['field'])): item['status'] for item in activity['attributes']}
                    expected_attributes = {(sku, evidence_tools().canonical_header(field)) for sku, fields in body['expected_rows'].items() for field in fields}
                    require(set(processing_attributes) == expected_attributes and len(processing_attributes) == len(activity['attributes']),
                            'processing_coverage', 'Processing status must cover every transmitted SKU and image field exactly once')
                observed = {(x.get('sku'), x.get('slot')): x for x in evidence.get('images', [])}
                processing_failures = []
                for image in body['images']:
                    item = observed.get((image['sku'], image['slot']), {})
                    field = ('other_product_image_locator_' + str(int(image['slot'][2:])) + '.0.media_location') if image['slot'].startswith('PT') else None
                    attribute_status = processing_attributes.get((image['sku'], field))
                    if attribute_status in {'failed', 'rejected'}:
                        processing_failures.append(f"{image['sku']}/{image['slot']}: {attribute_status}")
                        continue
                    # Amazon may transform files: collector must explicitly establish visual identity.
                    ok = item.get('source_sha256') == image['sha256'] and item.get('visually_verified') is True and item.get('live_url', '').startswith('https://') and attribute_status in {None, 'reflected'}
                    (matched if ok else pending).append(f"{image['sku']}/{image['slot']}")
            elif op == 'catalog.change':
                stage = evidence.get('stage')
                require(isinstance(stage, int) and 1 <= stage <= len(body['stages']), 'invalid_stage', 'Catalog evidence requires a valid stage')
                require(not state.get('execution_stage') or stage == state['execution_stage'], 'stage_mismatch', 'Evidence stage differs from current submission stage')
                previous = state.get('verified_stages', [])
                require(stage == 1 or stage - 1 in previous, 'stage_order', 'Verify the preceding stage before reconciling the next')
                for sku, fields in body['stages'][stage - 1].get('full_update_values', {}).items():
                    observed = evidence.get('catalog_rows', {}).get(sku, {})
                    for field, value in fields.items():
                        (matched if field in observed and as_text(observed[field]) == as_text(value) else pending).append(f'{sku}/{field}')
                manifest = body['manifest']
                family = manifest.get('family', {})
                offers = evidence.get('child_offers', {})
                for sku in body['protected_children']:
                    (matched if offers.get(sku) == 'preserved' else pending).append(f'{sku}/offer')
                if manifest['operation'] == 'create_products':
                    for product in body['products']:
                        live = evidence.get('products', {}).get(product['sku'], {})
                        ok = re.fullmatch(r'[A-Z0-9]{10}', str(live.get('asin', ''))) and live.get('product_type') == product['product_type'] and live.get('title') == product['title'] and live.get('parent_sku') in {'', None}
                        (matched if ok else pending).append(product['sku'])
                    status = 'verified' if matched and not pending and not evidence.get('errors') else ('partial' if matched else ('processing' if processing == 'live_observed' else 'failed'))
                    self._complete_reconciled_task(request, state, plan, status)
                    return self.persist(directory, state, status, matched=matched, pending=pending, evidence_path=str(evidence_path))
                parent = family['parent']['sku']
                deletion = len(body['stages']) > 1 and stage == 1 or manifest['operation'] == 'delete_parent'
                if deletion:
                    delete_sku = family.get('old_parent_sku') or parent
                    (matched if evidence.get('deleted_parents') == [delete_sku] else pending).append(f'{delete_sku}/deleted')
                    for sku in body['protected_children']:
                        (matched if evidence.get('relationships', {}).get(sku) == '' else pending).append(f'{sku}/detached')
                else:
                    for child in family.get('children', []):
                        sku = child['sku']
                        (matched if evidence.get('relationships', {}).get(sku) == parent else pending).append(f'{sku}/parent')
                if not pending and not evidence.get('errors'):
                    previous = sorted(set(previous + [stage]))
                    state['verified_stages'] = previous
                    if stage < len(body['stages']):
                        return self.persist(directory, state, 'partial', next_stage=stage + 1, evidence_path=str(evidence_path), message='Stage verified; next upload remains unsubmitted')
            elif op == 'shipment.create':
                expected = body['shipment']
                require(evidence.get('shipment_reference') == expected['shipment_reference'], 'shipment_mismatch', 'Shipment reference mismatch')
                require(evidence.get('carrier') == expected['carrier'] and evidence.get('currency') == expected['currency'], 'shipment_mismatch', 'Shipment carrier/currency mismatch')
                actual_cost = evidence.get('actual_cost')
                require(isinstance(actual_cost, (int, float)) and not isinstance(actual_cost, bool) and math.isfinite(actual_cost) and 0 <= actual_cost <= body['limits']['max_cost'], 'cost_limit', 'Actual shipment cost exceeds saved limit or is invalid')
                expected_qty = {x['sku']: x['quantity'] for x in expected['lines']}
                require(evidence.get('quantities') == expected_qty, 'shipment_mismatch', 'Confirmed quantities differ from shipment request')
                labels = evidence.get('labels', [])
                expected_boxes = {x['id'] for x in expected['cartons']}
                require(len(labels) == len(expected_boxes) and {x.get('carton_id') for x in labels} == expected_boxes, 'label_coverage', 'Every carton requires exactly one label record')
                shipment_ids = evidence.get('submission_ids', [evidence['submission_id']])
                require(isinstance(shipment_ids, list) and shipment_ids and len(set(shipment_ids)) == len(shipment_ids), 'invalid_label', 'Shipment IDs must be a unique nonempty list')
                verification_input = directory / 'label-verification-input.json'
                outputs = list({(x.get('path'), x.get('shipment_id')): {key: x.get(key) for key in ('path', 'sha256', 'shipment_id')} for x in labels}.values())
                atomic_json(verification_input, {'outputs': outputs, 'cartons': expected['cartons']})
                checked = subprocess.run(['node', str(ROOT / 'tools/amazon-operations/verify-shipment-labels.mjs'), '--request', str(verification_input)], capture_output=True, text=True, timeout=60, check=False)
                checked_result = json.loads(checked.stdout)
                require(checked.returncode == 0 and checked_result.get('status') == 'verified', 'invalid_label', 'Independent PDF content/page/thermal-size verification failed')
                checked_labels = {x['carton_id']: x for x in checked_result['labels']}
                require(set(checked_labels) == expected_boxes, 'label_coverage', 'Independent PDF carton coverage differs')
                for label in labels:
                    require(all(checked_labels[label['carton_id']].get(key) == label.get(key) for key in ('path', 'sha256', 'shipment_id')), 'invalid_label', 'Declared label identity differs from parsed PDF content')
                    path = Path(label.get('path', ''))
                    require(path.is_file() and file_hash(path) == label.get('sha256'), 'invalid_label', 'Label file checksum mismatch')
                    with path.open('rb') as stream:
                        require(stream.read(5) == b'%PDF-', 'invalid_label', 'Shipment labels must be PDF files')
                    require(label.get('shipment_id') in shipment_ids, 'invalid_label', 'Label shipment identity mismatch')
                    matched.append(label['carton_id'])
            errors = evidence.get('errors', []) + (processing_failures if op == 'listing.images' else [])
            status = 'verified' if matched and not pending and not errors else ('partial' if matched else ('processing' if processing == 'live_observed' else 'failed'))
            if op == 'listing.images' and errors and not pending:
                # Retain the verified subset while stopping a fully resolved
                # rejected batch; polling cannot repair terminal feed failures.
                status = 'failed'
            self._complete_reconciled_task(request, state, plan, status)
            return self.persist(directory, state, status, submission_id=evidence.get('submission_id'), matched=matched, pending=pending, failures=errors, evidence_path=str(evidence_path))


def capabilities():
    return {'schema_version': 1, 'operations': {name: {'prepare': True, 'reconcile': True, 'execute_adapter': {'seo.update':'flatfilepro.cdp','flatfilepro.update':'flatfilepro.cdp','catalog.change':'catalog.cdp','shipment.create':'shipment.cdp','listing.images':'flatfilepro.cdp','case.create':'cases.cdp','case.reply':'cases.cdp'}.get(operation_kind(name)), 'production_ready': False, 'requires_live_canary': True, 'limitation': 'Live selector contracts require scoped canary' if operation_kind(name) in {'seo.update', 'flatfilepro.update', 'catalog.change', 'shipment.create', 'listing.images', 'case.create', 'case.reply'} else 'No observed submission adapter installed'} for name in sorted(OPERATIONS)}}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=['capabilities', 'prepare', 'execute', 'advance', 'reconcile', 'status', 'observe', 'record-image-review', 'image-release-proof'])
    parser.add_argument('--request', type=Path)
    parser.add_argument('--state-dir', type=Path)
    args = parser.parse_args(argv)
    try:
        if args.command == 'capabilities':
            result = capabilities()
        else:
            require(args.request is not None and args.state_dir is not None, 'missing_argument', '--request and --state-dir are required')
            request = json.loads(args.request.read_text())
            service = Operations(args.state_dir)
            command = 'execute' if args.command == 'advance' else args.command.replace('-', '_')
            result = service.view(service.directory(request.get('operation_id'))) if command == 'status' else getattr(service, command)(request)
        print(json.dumps(result, ensure_ascii=False, allow_nan=False))
        return 0
    except (OperationError, OSError, ValueError, KeyError, TypeError, ImportError, SystemExit) as exc:
        print(json.dumps({'schema_version': 1, 'status': 'blocked', 'reason': getattr(exc, 'code', 'invalid_input'), 'message': str(exc)}))
        return 2


if __name__ == '__main__':
    sys.exit(main())
