#!/usr/bin/env python3
"""Synthetic, client-neutral tests for prepare_flatfilepro_upload.py.

Run: python3 skills/amazon-flatfilepro/scripts/selftest_prepare.py
(openpyxl is required). Prints a PASS/FAIL summary and exits non-zero on failure.
"""
from __future__ import annotations

import contextlib
import csv
import io
import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import prepare_flatfilepro_upload as prep  # noqa: E402


TMP = Path(tempfile.mkdtemp(prefix="ffp-prepare-selftest-"))
FAILURES: list[str] = []
MANGLED = re.compile(r"__\d+__")


def check(label: str, condition: bool, detail: str = "") -> None:
    status = "ok" if condition else "FAIL"
    print(f"  [{status}] {label}" + (f": {detail}" if detail and not condition else ""))
    if not condition:
        FAILURES.append(f"{label}: {detail}")


def write_csv(path: Path, rows: list[list[str]]) -> Path:
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)
    return path


def write_xlsx(path: Path, rows: list[list[str]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()
    return path


def write_changes(path: Path, changes: list[tuple[str, str, str]]) -> Path:
    return write_csv(path, [["sku", "attribute", "value"], *[list(c) for c in changes]])


def run(source: Path, changes: Path, output: Path, *extra: str) -> tuple[int, str, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        code = prep.main(["--source", str(source), "--changes", str(changes), "--output", str(output), *extra])
    return code, out.getvalue(), err.getvalue()


def read_output(path: Path) -> tuple[list[str], dict[str, dict[str, str]]]:
    if path.suffix.lower() == ".xlsx":
        wb = openpyxl.load_workbook(path)
        rows = [["" if v is None else str(v) for v in row] for row in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        with path.open(newline="", encoding="utf-8") as f:
            rows = list(csv.reader(f))
    headers = rows[0]
    return headers, {r[0]: dict(zip(headers, r)) for r in rows[1:]}


OUTPUTS: list[Path] = []


def case_canonical_csv() -> None:
    print("\n=== canonical CSV source ===")
    src = write_csv(TMP / "canonical.csv", [
        ["sku", "item_name.0.value", "bullet_point.0.value", "bullet_point.1.value"],
        ["TEST-SKU-001", "Old title one", "Old bullet A", "Old bullet B"],
        ["TEST-SKU-002", "Old title two", "Old bullet C", "Old bullet D"],
    ])
    chg = write_changes(TMP / "canonical-changes.csv", [("TEST-SKU-001", "item_name.0.value", "New title one")])
    out = TMP / "canonical-out.xlsx"
    code, _, err = run(src, chg, out)
    check("canonical source exits 0", code == 0, err)
    headers, rows = read_output(out)
    OUTPUTS.append(out)
    check("canonical headers are sku plus the changed field", headers == ["sku", "item_name.0.value"], str(headers))
    check("canonical change written", rows.get("TEST-SKU-001", {}).get("item_name.0.value") == "New title one", str(rows))
    check("only the changed SKU is written", list(rows) == ["TEST-SKU-001"], str(list(rows)))


def case_mangled_xlsx() -> None:
    print("\n=== mangled __N__ XLSX source ===")
    src = write_xlsx(TMP / "mangled.xlsx", [
        ["FlatFilePro export (synthetic)"],
        ["sku", "asin", "item_name__1__value", "unit_count__1__value", "unit_count__1__type_value", "bullet_point__1__value"],
        ["TEST-SKU-101", "B0TEST0101", "Widget title", "60", "Count", "Bullet one"],
        ["TEST-SKU-102", "B0TEST0102", "Gadget title", "12", "Count", "Bullet two"],
    ])
    headers = prep.read_source_headers(src)
    check("header row located below a preamble row", headers[0] == "sku" and "unit_count__1__type_value" in headers, str(headers))
    check("nested header canonicalises to unit_count.0.type_value",
          prep.canonical_attribute("unit_count__1__type_value") == "unit_count.0.type_value",
          prep.canonical_attribute("unit_count__1__type_value"))
    chg = write_changes(TMP / "mangled-changes.csv", [
        ("TEST-SKU-101", "unit_count__1__value", "30"),
        ("TEST-SKU-101", "unit_count__1__type_value", "Count"),
        ("TEST-SKU-102", "item_name__1__value", "Gadget title, revised"),
    ])
    out = TMP / "mangled-out.xlsx"
    code, _, err = run(src, chg, out)
    check("mangled source exits 0", code == 0, err)
    out_headers, rows = read_output(out)
    OUTPUTS.append(out)
    check("mangled headers written in canonical source order",
          out_headers == ["sku", "item_name.0.value", "unit_count.0.value", "unit_count.0.type_value"], str(out_headers))
    check("nested change written under canonical header",
          rows.get("TEST-SKU-101", {}).get("unit_count.0.type_value") == "Count"
          and rows["TEST-SKU-101"].get("unit_count.0.value") == "30", str(rows))
    check("unchanged cells stay empty without --fill-unchanged",
          rows.get("TEST-SKU-102", {}).get("unit_count.0.value") == "" and rows["TEST-SKU-101"].get("item_name.0.value") == "",
          str(rows))


EXPORT_HEADERS = [
    "sku", "item_name__1__value", "bullet_point__1__value", "bullet_point__2__value",
    "color__1__standardized_values__1__", "purchasable_offer__1__our_price__1__schedule__1__value_with_tax",
    "unit_count__1__type_value",
]


def case_export_spelled_csv() -> None:
    print("\n=== FFP-export-spelled CSV source (MCP grid shape) ===")
    src = write_csv(TMP / "export-grid.csv", [
        EXPORT_HEADERS,
        ["TEST-SKU-201", "Blue widget", "Bullet 1", "Bullet 2", "Blue", "49.99", "Count"],
        ["TEST-SKU-202", "Red widget", "Bullet 3", "Bullet 4", "Red", "39.99", "Count"],
    ])
    chg = write_changes(TMP / "export-changes.csv", [
        ("TEST-SKU-201", "bullet_point__2__value", "Bullet 2, revised"),
        ("TEST-SKU-202", "item_name__1__value", "Red widget, revised"),
    ])
    out = TMP / "export-out.csv"
    code, _, err = run(src, chg, out, "--fill-unchanged")
    check("export-spelled source exits 0", code == 0, err)
    headers, rows = read_output(out)
    OUTPUTS.append(out)
    check("export-spelled headers canonicalised",
          headers == ["sku", "item_name.0.value", "bullet_point.1.value"], str(headers))
    check("change and carried value both present",
          rows.get("TEST-SKU-201", {}).get("bullet_point.1.value") == "Bullet 2, revised"
          and rows["TEST-SKU-201"].get("item_name.0.value") == "Blue widget", str(rows))
    check("array-of-strings and nested-array headers canonicalise",
          prep.canonical_attribute("color__1__standardized_values__1__") == "color.0.standardized_values.0."
          and prep.canonical_attribute("purchasable_offer__1__our_price__1__schedule__1__value_with_tax")
          == "purchasable_offer.0.our_price.0.schedule.0.value_with_tax")


def case_missing_change_sku() -> None:
    print("\n=== changed SKU missing from source ===")
    src = write_csv(TMP / "missing.csv", [
        ["sku", "item_name__1__value"],
        ["TEST-SKU-301", "Title"],
    ])
    chg = write_changes(TMP / "missing-changes.csv", [
        ("TEST-SKU-301", "item_name__1__value", "New"),
        ("TEST-SKU-399", "item_name__1__value", "New"),
        ("-TEST-SKU-398", "item_name__1__value", "New"),
    ])
    out = TMP / "missing-out.xlsx"
    code, _, err = run(src, chg, out)
    check("missing change SKU exits 2", code == 2, f"exit {code}")
    check("every missing SKU is named", "TEST-SKU-399" in err and "-TEST-SKU-398" in err, err)
    check("present SKU is not reported as missing", "TEST-SKU-301" not in err, err)
    check("no output written", not out.exists())


def case_duplicate_source_sku() -> None:
    print("\n=== duplicate source SKU rows ===")
    src = write_xlsx(TMP / "duplicate.xlsx", [
        ["sku", "item_name__1__value"],
        ["TEST-SKU-401", "Title A"],
        ["TEST-SKU-402", "Title B"],
        ["TEST-SKU-401", "Title A again"],
        ["TEST-SKU-403", "Title C"],
        ["TEST-SKU-402", "Title B again"],
        ["TEST-SKU-402", "Title B third"],
    ])
    chg = write_changes(TMP / "duplicate-changes.csv", [("TEST-SKU-403", "item_name__1__value", "New")])
    out = TMP / "duplicate-out.xlsx"
    code, _, err = run(src, chg, out, "--fill-unchanged")
    check("duplicate source SKU exits 2", code == 2, f"exit {code}")
    check("every duplicate SKU is named with its row count",
          "TEST-SKU-401 (2 rows)" in err and "TEST-SKU-402 (3 rows)" in err, err)
    check("unique SKU is not reported", "TEST-SKU-403" not in err, err)
    check("no output written", not out.exists())

    both = write_changes(TMP / "both-changes.csv", [("TEST-SKU-499", "item_name__1__value", "New")])
    code, _, err = run(src, both, TMP / "both-out.xlsx")
    check("missing and duplicate SKUs reported together",
          code == 2 and "TEST-SKU-499" in err and "TEST-SKU-401" in err, err)


def case_fill_unchanged_full_grid() -> None:
    print("\n=== --fill-unchanged full grid ===")
    src = write_xlsx(TMP / "grid.xlsx", [
        EXPORT_HEADERS,
        ["TEST-SKU-501", "Title 1", "B1-1", "B1-2", "Blue", "10.00", "Count"],
        ["TEST-SKU-502", "Title 2", "B2-1", "B2-2", "Red", "20.00", "Count"],
        ["TEST-SKU-503", "Title 3", "B3-1", "B3-2", "Green", "30.00", "Ounce"],
    ])
    chg = write_changes(TMP / "grid-changes.csv", [
        ("TEST-SKU-501", "item_name__1__value", "Title 1 new"),
        ("TEST-SKU-502", "bullet_point__1__value", "B2-1 new"),
        ("TEST-SKU-503", "unit_count__1__type_value", "Count"),
    ])
    out = TMP / "grid-out.xlsx"
    code, stdout, err = run(src, chg, out, "--fill-unchanged")
    check("full grid exits 0", code == 0, err)
    check("full grid mode is announced", "full-grid" in stdout, stdout)
    headers, rows = read_output(out)
    OUTPUTS.append(out)
    expected = {
        "TEST-SKU-501": {"sku": "TEST-SKU-501", "item_name.0.value": "Title 1 new", "bullet_point.0.value": "B1-1", "unit_count.0.type_value": "Count"},
        "TEST-SKU-502": {"sku": "TEST-SKU-502", "item_name.0.value": "Title 2", "bullet_point.0.value": "B2-1 new", "unit_count.0.type_value": "Count"},
        "TEST-SKU-503": {"sku": "TEST-SKU-503", "item_name.0.value": "Title 3", "bullet_point.0.value": "B3-1", "unit_count.0.type_value": "Count"},
    }
    check("full grid headers", headers == ["sku", "item_name.0.value", "bullet_point.0.value", "unit_count.0.type_value"], str(headers))
    check("every cell is a change or the source value", rows == expected, str(rows))
    check("row count matches changed SKUs", len(rows) == 3, str(len(rows)))


def case_no_mangled_output_headers() -> None:
    print("\n=== output headers ===")
    bad = {str(p.name): [h for h in read_output(p)[0] if MANGLED.search(h)] for p in OUTPUTS}
    bad = {k: v for k, v in bad.items() if v}
    check(f"no __N__ header in {len(OUTPUTS)} outputs", bool(OUTPUTS) and not bad, str(bad))


def main() -> int:
    print(f"Working dir: {TMP}")
    for case in (case_canonical_csv, case_mangled_xlsx, case_export_spelled_csv, case_missing_change_sku,
                 case_duplicate_source_sku, case_fill_unchanged_full_grid, case_no_mangled_output_headers):
        try:
            case()
        except Exception as exc:  # report and continue so the summary lists every failure
            check(f"{case.__name__} raised", False, repr(exc))

    print("\n=== Result ===")
    if FAILURES:
        print(f"FAIL ({len(FAILURES)})")
        for failure in FAILURES:
            print(f"- {failure}")
        return 1
    print("PASS")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    finally:
        if not FAILURES:
            shutil.rmtree(TMP, ignore_errors=True)
